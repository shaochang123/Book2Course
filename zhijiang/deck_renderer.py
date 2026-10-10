"""Template-aware native PPT pages, reusing verified SVG and narrated clips."""
from __future__ import annotations
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from zhijiang.presentation_templates import get_template
from zhijiang.presentation import (
    PresentationError, _color, _font, _safe_text,
)

PAGE_W, PAGE_H = 13.333, 7.5
ROLES = {'explanation': '概念解释', 'example': '示例讲解', 'comparison': '对比观察',
         'evidence': '原文与依据', 'recap': '回顾理解'}


def _fit_text(text, width, height, max_size=26, min_size=16):
    """Same wrapping/size choice for editable PPT text and the preview."""
    probe = ImageDraw.Draw(Image.new('RGB', (1, 1)))
    min_size = min(min_size, max_size)
    for size in range(max_size, min_size-1, -1):
        font = _font(round(size*96/72))
        lines, current = [], ''
        for char in _safe_text(text):
            if char == '\n':
                lines.append(current); current = ''; continue
            if current and probe.textlength(current+char, font=font) > width*96:
                lines.append(current); current = ''
            current += char
        if current:
            lines.append(current)
        if len(lines)*font.size*1.25 <= height*96:
            return '\n'.join(lines), size
    raise PresentationError('文字超出画面，请换版式或拆页。')


class Page:
    """Native shape/text commands also produce a local review PNG."""
    def __init__(self, deck, template, folder, number):
        self.slide = deck.slides.add_slide(deck.slide_layouts[6])
        self.template = template
        self.preview = Image.new('RGB', (1280, 720), '#'+template.background)
        self.draw = ImageDraw.Draw(self.preview)
        self.folder, self.number = folder, number
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = _color(template.background)
        if template.id == 'paper':
            # Original ruled-paper accent, no reference artwork.
            for y in (1.3, 6.95):
                self.rect((.55, y, 12.23, .018), template.panel)
            self.rect((.55, .35, .12, .7), template.accent)
        elif template.id == 'academic':
            self.rect((0, 0, PAGE_W, .16), template.accent)
            self.rect((.4, .45, .06, 6.35), template.accent)
        elif template.id == 'editorial':
            self.rect((.55, .4, .45, 6.5), template.accent)

    def rect(self, box, color, *, rounded=False):
        x, y, w, h = box
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = _color(color)
        shape.line.fill.background()
        bounds = tuple(round(v*96) for v in (x, y, x+w, y+h))
        if rounded:
            self.draw.rounded_rectangle(bounds, 14, fill='#'+color)
        else:
            self.draw.rectangle(bounds, fill='#'+color)

    def text(self, value, box, *, size=24, min_size=24, color=None, bold=False):
        x, y, w, h = box
        value, size = _fit_text(value, w, h, size, min_size)
        shape = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame; tf.clear(); tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(value.splitlines()):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.25
            p.space_after = Pt(0)
            run = p.add_run(); run.text = line
            run.font.name = 'Microsoft YaHei'; run.font.size = Pt(size)
            run.font.bold = bold; run.font.color.rgb = _color(color or self.template.ink)
        font = _font(round(size*96/72), bold=bold)
        for i, line in enumerate(value.splitlines()):
            self.draw.text((x*96, y*96+i*font.size*1.25), line,
                           fill='#'+(color or self.template.ink), font=font)

    def picture(self, path, box):
        x, y, w, h = fit_picture(path, box)
        shape = self.slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        with Image.open(path) as im:
            im = im.convert('RGBA'); im = im.resize((round(w*96), round(h*96)), Image.Resampling.LANCZOS)
            self.preview.paste(im, (round(x*96), round(y*96)), im)
        return shape

    def footer(self, source=''):
        self.text(source, (1.3 if self.template.id == 'editorial' else .8, 7.08, 10.2, .23),
                  size=11, min_size=10, color=self.template.muted)
        self.text(f'{self.number:02d}', (12.1, 7.05, .6, .3), size=12,
                  color=self.template.muted)

    def save_preview(self):
        path = self.folder / f'page-{self.number:03d}.png'
        self.preview.save(path)
        return path


def fit_picture(path, box):
    x, y, width, height = box
    with Image.open(path) as image:
        ratio = image.width / image.height
    width = min(width, height*ratio)
    height = width/ratio
    return (x+(box[2]-width)/2, y+(box[3]-height)/2, width, height)


def _header(page, title, subtitle):
    left = 1.3 if page.template.id == 'editorial' else .8
    page.text(title, (left, .35, 11.0, .83), size=page.template.title_size,
              min_size=20, bold=True)
    page.text(subtitle, (left, 1.14, 11.0, .28), size=13,
              min_size=11, color=page.template.muted)


def render_template_presentation(lesson, output, *, audio_files=None):
    from zhijiang.rich_deck import render_shared_deck
    return render_shared_deck(lesson, output, audio_files=audio_files)
