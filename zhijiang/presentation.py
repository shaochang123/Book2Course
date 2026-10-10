"""Render a source-backed lesson as a PPTX with SVG diagrams and playable clips.

python-pptx creates ordinary slides and video shapes. It cannot insert SVG pictures
directly, so each diagram gets a PNG fallback plus the Office SVG blip extension.
The original SVG bytes remain in the PPTX media folder.
"""

from __future__ import annotations

import html
import os
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

import imageio_ffmpeg
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from zhijiang.models import Lesson, LessonSegment


INK = "081623"
PANEL = "10283A"
PANEL_LIGHT = "18364B"
ACCENT = "6DE2C0"
WHITE = "F4F7F9"
MUTED = "A7C0CC"
LINE = "345569"

SVG_NS = "http://www.w3.org/2000/svg"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
REL_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


class PresentationError(RuntimeError):
    """A teaching presentation could not be rendered."""


def _color(hex_rgb: str) -> RGBColor:
    return RGBColor.from_string(hex_rgb)


def _safe_text(value: str) -> str:
    # XML 1.0 permits tab/newline/CR and ordinary printable characters.
    return "".join(char for char in value if char in "\t\n\r" or ord(char) >= 32).strip()


def _xml_text(value: str) -> str:
    return html.escape(_safe_text(value), quote=True)


def _display_lines(value: str, width: int = 18, max_lines: int = 3) -> list[str]:
    """Bound visual text while keeping full wording in the speaker notes."""
    value = re.sub(r"\s+", " ", _safe_text(value))
    lines: list[str] = []
    line = ""
    units = 0
    for char in value:
        size = 1 if ord(char) > 0x2E80 else 0.55
        if units + size > width and line:
            lines.append(line.strip())
            line, units = "", 0
            if len(lines) == max_lines:
                break
        line += char
        units += size
    if len(lines) < max_lines and line:
        lines.append(line.strip())
    if len("".join(lines).replace(" ", "")) < len(value.replace(" ", "")) and lines:
        lines[-1] = lines[-1].rstrip("… ") + "…"
    return lines or [""]


def _svg_text(x: float, y: float, text: str, *, size: int = 27,
              color: str = WHITE, weight: int = 400, anchor: str = "middle") -> str:
    return (f'<text x="{x:.1f}" y="{y:.1f}" text-anchor="{anchor}" '
            f'font-family="Microsoft YaHei, Noto Sans CJK SC, sans-serif" '
            f'font-size="{size}" font-weight="{weight}" fill="#{color}">'
            f'{_xml_text(text)}</text>')


def _svg_multiline(x: float, top: float, value: str, *, width: int,
                   max_lines: int = 3, size: int = 26) -> str:
    lines = _display_lines(value, width, max_lines)
    return "".join(_svg_text(x, top + index * (size + 12), line, size=size)
                   for index, line in enumerate(lines))


def _diagram_svg(segment: LessonSegment) -> bytes:
    """Source-grounded diagram; relationships only reflect segment.kind."""
    bullets = segment.bullets[:4]
    parts = [
        f'<svg xmlns="{SVG_NS}" width="1200" height="450" viewBox="0 0 1200 450" '
        'role="img">',
        f'<title>{_xml_text(segment.title)}</title>',
        f'<desc>{_xml_text("；".join(bullets))}</desc>',
        f'<rect width="1200" height="450" rx="24" fill="#{PANEL}"/>',
    ]
    if segment.kind == "process":
        n = len(bullets)
        gap = 24
        width = min(270, (1080 - gap * (n - 1)) / n)
        left = (1200 - (n * width + (n - 1) * gap)) / 2
        for index, bullet in enumerate(bullets):
            x = left + index * (width + gap)
            if index:
                parts.append(f'<path d="M {x-gap+4:.1f} 220 L {x-6:.1f} 220" '
                             f'stroke="#{ACCENT}" stroke-width="4"/>'
                             f'<path d="M {x-13:.1f} 211 L {x-5:.1f} 220 L {x-13:.1f} 229" '
                             f'fill="none" stroke="#{ACCENT}" stroke-width="4"/>')
            parts.append(f'<rect x="{x:.1f}" y="90" width="{width:.1f}" height="278" '
                         f'rx="22" fill="#{PANEL_LIGHT}" stroke="#{LINE}" stroke-width="2"/>')
            parts.append(f'<circle cx="{x+width/2:.1f}" cy="140" r="25" fill="#{ACCENT}"/>')
            parts.append(_svg_text(x + width / 2, 149, str(index + 1), size=25,
                                   color=INK, weight=700))
            parts.append(_svg_multiline(x + width / 2, 202, bullet,
                                        width=12 if n >= 4 else 16, max_lines=4,
                                        size=25 if n >= 4 else 28))
    elif segment.kind == "formula":
        n = len(bullets)
        row_height = min(88, 316 / n)
        top = (450 - n * row_height - (n - 1) * 14) / 2
        for index, bullet in enumerate(bullets):
            y = top + index * (row_height + 14)
            parts.append(f'<rect x="132" y="{y:.1f}" width="936" height="{row_height:.1f}" '
                         f'rx="17" fill="#{PANEL_LIGHT}" stroke="#{LINE}" stroke-width="2"/>')
            parts.append(f'<circle cx="178" cy="{y+row_height/2:.1f}" r="23" '
                         f'fill="#{ACCENT}"/>')
            parts.append(_svg_text(178, y + row_height / 2 + 9, str(index + 1),
                                   size=24, color=INK, weight=700))
            lines = _display_lines(bullet, 50, 2)
            for line_index, line in enumerate(lines):
                parts.append(_svg_text(225, y + row_height / 2 + 9 +
                                       (line_index - (len(lines)-1)/2) * 30,
                                       line, size=26, anchor="start"))
    else:  # concept: spokes connect related points without implying causality.
        n = len(bullets)
        places = {1: [(600, 337)], 2: [(330, 326), (870, 326)],
                  3: [(250, 325), (600, 325), (950, 325)],
                  4: [(175, 322), (455, 322), (745, 322), (1025, 322)]}[n]
        parts.append(f'<rect x="382" y="52" width="436" height="114" rx="28" '
                     f'fill="#{PANEL_LIGHT}" stroke="#{ACCENT}" stroke-width="3"/>')
        parts.append(_svg_multiline(600, 103, segment.title, width=22,
                                    max_lines=2, size=30))
        for index, (cx, cy) in enumerate(places):
            parts.append(f'<path d="M 600 166 Q 600 240 {cx} 253" fill="none" '
                         f'stroke="#{LINE}" stroke-width="3"/>')
            width = 254 if n == 4 else 320
            x = cx - width / 2
            parts.append(f'<rect x="{x:.1f}" y="260" width="{width}" height="145" '
                         f'rx="20" fill="#{PANEL_LIGHT}" stroke="#{LINE}" stroke-width="2"/>')
            parts.append(f'<circle cx="{cx}" cy="273" r="11" fill="#{ACCENT}"/>')
            parts.append(_svg_multiline(cx, 317, bullets[index], width=11 if n == 4 else 16,
                                        max_lines=3, size=24 if n == 4 else 27))
    parts.append("</svg>")
    result = "".join(parts).encode("utf-8")
    etree.fromstring(result)  # Fail early on malformed XML or unsafe text.
    return result


def _font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    paths = (
        ["C:/Windows/Fonts/msyhbd.ttc", "C:/Windows/Fonts/simhei.ttf"] if bold
        else ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf"]
    ) + [
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in paths:
        if Path(path).is_file():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default(size=size)


def _pil_lines(draw: ImageDraw.ImageDraw, value: str, max_width: int,
               font: ImageFont.FreeTypeFont, max_lines: int = 3) -> list[str]:
    chars = re.sub(r"\s+", " ", _safe_text(value))
    lines: list[str] = []
    line = ""
    for char in chars:
        if draw.textlength(line + char, font=font) > max_width and line:
            lines.append(line)
            line = ""
            if len(lines) == max_lines:
                break
        line += char
    if line and len(lines) < max_lines:
        lines.append(line)
    if len("".join(lines)) < len(chars) and lines:
        while lines[-1] and draw.textlength(lines[-1] + "…", font=font) > max_width:
            lines[-1] = lines[-1][:-1]
        lines[-1] += "…"
    return lines or [""]


def _draw_centered(draw: ImageDraw.ImageDraw, text: str, center_x: float,
                   top: float, font: ImageFont.FreeTypeFont, color: str,
                   max_width: int, max_lines: int = 3, gap: int = 6) -> None:
    for index, line in enumerate(_pil_lines(draw, text, max_width, font, max_lines)):
        box = draw.textbbox((0, 0), line, font=font)
        draw.text((center_x - (box[2] - box[0]) / 2, top + index * (font.size + gap)),
                  line, font=font, fill=color)


def _diagram_png(segment: LessonSegment, output: Path) -> None:
    """Raster fallback mirrors the SVG's information for older PowerPoint clients."""
    image = Image.new("RGB", (1200, 450), f"#{PANEL}")
    draw = ImageDraw.Draw(image)
    bullets = segment.bullets[:4]
    if segment.kind == "process":
        n = len(bullets)
        gap = 24
        width = min(270, (1080 - gap * (n - 1)) / n)
        left = (1200 - (n * width + (n - 1) * gap)) / 2
        for index, bullet in enumerate(bullets):
            x = left + index * (width + gap)
            if index:
                draw.line((x-gap+4, 220, x-6, 220), fill=f"#{ACCENT}", width=4)
                draw.polygon([(x-13, 211), (x-5, 220), (x-13, 229)], fill=f"#{ACCENT}")
            draw.rounded_rectangle((x, 90, x+width, 368), radius=22,
                                   fill=f"#{PANEL_LIGHT}", outline=f"#{LINE}", width=2)
            cx = x + width / 2
            draw.ellipse((cx-25, 115, cx+25, 165), fill=f"#{ACCENT}")
            _draw_centered(draw, str(index+1), cx, 121, _font(25, bold=True),
                           f"#{INK}", 45, 1)
            _draw_centered(draw, bullet, cx, 192, _font(25 if n >= 4 else 28),
                           f"#{WHITE}", int(width)-35, 4)
    elif segment.kind == "formula":
        n = len(bullets)
        h = min(88, 316 / n)
        top = (450 - n * h - (n-1)*14) / 2
        for index, bullet in enumerate(bullets):
            y = top + index * (h+14)
            draw.rounded_rectangle((132, y, 1068, y+h), radius=17,
                                   fill=f"#{PANEL_LIGHT}", outline=f"#{LINE}", width=2)
            draw.ellipse((155, y+h/2-23, 201, y+h/2+23), fill=f"#{ACCENT}")
            _draw_centered(draw, str(index+1), 178, y+h/2-15,
                           _font(24, bold=True), f"#{INK}", 40, 1)
            lines = _pil_lines(draw, bullet, 805, _font(26), 2)
            for j, line in enumerate(lines):
                draw.text((225, y+h/2-15+(j-(len(lines)-1)/2)*30), line,
                          font=_font(26), fill=f"#{WHITE}")
    else:
        n = len(bullets)
        places = {1: [(600, 337)], 2: [(330, 326), (870, 326)],
                  3: [(250, 325), (600, 325), (950, 325)],
                  4: [(175, 322), (455, 322), (745, 322), (1025, 322)]}[n]
        draw.rounded_rectangle((382, 52, 818, 166), radius=28,
                               fill=f"#{PANEL_LIGHT}", outline=f"#{ACCENT}", width=3)
        _draw_centered(draw, segment.title, 600, 88, _font(30, bold=True),
                       f"#{WHITE}", 400, 2)
        for index, (cx, _) in enumerate(places):
            draw.line((600, 166, cx, 260), fill=f"#{LINE}", width=3)
            w = 254 if n == 4 else 320
            draw.rounded_rectangle((cx-w/2, 260, cx+w/2, 405), radius=20,
                                   fill=f"#{PANEL_LIGHT}", outline=f"#{LINE}", width=2)
            draw.ellipse((cx-11, 262, cx+11, 284), fill=f"#{ACCENT}")
            _draw_centered(draw, bullets[index], cx, 304,
                           _font(24 if n == 4 else 27), f"#{WHITE}", w-32, 3)
    image.save(output)


def _ease(value: float) -> float:
    value = min(1.0, max(0.0, value))
    return value * value * (3 - 2 * value)


def _animation_frame(segment: LessonSegment, index: int, total: int, progress: float) -> Image.Image:
    """Animate the actual concept/step relationships rather than slide decoration."""
    image = Image.new("RGB", (960, 540), f"#{INK}")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((28, 28, 932, 512), radius=26, fill=f"#{PANEL}")
    draw.text((54, 43), f"{index:02d} / {total:02d}   ·   {segment.kind.upper()}",
              font=_font(21, bold=True), fill=f"#{ACCENT}")
    title = _pil_lines(draw, segment.title, 845, _font(35, bold=True), 1)[0]
    draw.text((54, 89), title, font=_font(35, bold=True), fill=f"#{WHITE}")
    bullets = segment.bullets[:4]
    n = len(bullets)
    stage = min(n - 1, int(progress * n))
    phase = _ease(progress * n - stage)
    if segment.kind == "process":
        gap = 20
        w = min(206, (836 - gap*(n-1))/n)
        x0 = (960 - (n*w+(n-1)*gap))/2
        y = 235
        full_line_left = x0 + w/2
        full_line_right = x0 + (n-1)*(w+gap) + w/2
        end = full_line_left + (full_line_right-full_line_left) * progress
        draw.line((full_line_left, y, full_line_right, y), fill=f"#{LINE}", width=6)
        draw.line((full_line_left, y, end, y), fill=f"#{ACCENT}", width=8)
        for j, bullet in enumerate(bullets):
            x = x0 + j*(w+gap)
            active = j < stage or (j == stage and phase > 0.14)
            draw.rounded_rectangle((x, 168, x+w, 438), radius=21,
                                   fill=f"#{PANEL_LIGHT if active else PANEL}",
                                   outline=f"#{ACCENT if j == stage else LINE}", width=3)
            draw.ellipse((x+w/2-27, y-27, x+w/2+27, y+27),
                         fill=f"#{ACCENT if active else LINE}")
            _draw_centered(draw, str(j+1), x+w/2, y-17, _font(27, bold=True),
                           f"#{INK if active else WHITE}", 40, 1)
            if active:
                _draw_centered(draw, bullet, x+w/2, 292, _font(24),
                               f"#{WHITE}", int(w-24), 4)
    elif segment.kind == "formula":
        h = min(69, 270/n)
        y0 = 168
        end = y0 + (n-1)*(h+15)+h/2
        draw.line((92, y0+h/2, 92, end), fill=f"#{LINE}", width=5)
        draw.line((92, y0+h/2, 92, y0+h/2+(end-y0-h/2)*progress),
                  fill=f"#{ACCENT}", width=7)
        for j, bullet in enumerate(bullets):
            y = y0+j*(h+15)
            active = j < stage or (j == stage and phase > 0.14)
            draw.rounded_rectangle((130, y, 884, y+h), radius=17,
                                   fill=f"#{PANEL_LIGHT if active else PANEL}",
                                   outline=f"#{ACCENT if j == stage else LINE}", width=2)
            draw.ellipse((73, y+h/2-19, 111, y+h/2+19),
                         fill=f"#{ACCENT if active else LINE}")
            _draw_centered(draw, str(j+1), 92, y+h/2-14,
                           _font(22, bold=True), f"#{INK if active else WHITE}", 30, 1)
            if active:
                line = _pil_lines(draw, bullet, 700, _font(25), 1)[0]
                draw.text((154, y+h/2-16), line, font=_font(25), fill=f"#{WHITE}")
    else:
        draw.rounded_rectangle((300, 160, 660, 245), radius=23,
                               fill=f"#{PANEL_LIGHT}", outline=f"#{ACCENT}", width=3)
        _draw_centered(draw, segment.title, 480, 181,
                       _font(27, bold=True), f"#{WHITE}", 330, 2)
        gap = 16
        w = min(206, (838-gap*(n-1))/n)
        x0 = (960-(n*w+(n-1)*gap))/2
        for j, bullet in enumerate(bullets):
            x = x0+j*(w+gap)
            cx = x+w/2
            y = 315
            active = j < stage or (j == stage and phase > 0.14)
            if active:
                target_y = 258 + (y-258)*(phase if j == stage else 1)
                draw.line((480, 245, cx, target_y), fill=f"#{ACCENT}", width=4)
            else:
                draw.line((480, 245, cx, y), fill=f"#{LINE}", width=2)
            draw.rounded_rectangle((x, y, x+w, 453), radius=18,
                                   fill=f"#{PANEL_LIGHT if active else PANEL}",
                                   outline=f"#{ACCENT if j == stage else LINE}", width=2)
            if active:
                _draw_centered(draw, bullet, cx, y+24, _font(23),
                               f"#{WHITE}", int(w-24), 3)
    draw.text((55, 480), f"来源：PDF 第 {segment.evidence.page} 页",
              font=_font(19), fill=f"#{MUTED}")
    return image


def _render_animation(segment: LessonSegment, number: int, total: int, output: Path,
                      poster: Path) -> None:
    fps, seconds = 12, 5.0
    frame_count = int(fps * seconds)
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    command = [ffmpeg, "-hide_banner", "-loglevel", "error", "-y",
               "-f", "rawvideo", "-pixel_format", "rgb24", "-video_size", "960x540",
               "-framerate", str(fps), "-i", "pipe:0", "-an", "-c:v", "libx264",
               "-pix_fmt", "yuv420p", "-crf", "25", "-movflags", "+faststart",
               str(output)]
    process = subprocess.Popen(command, stdin=subprocess.PIPE,
                               stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    try:
        assert process.stdin is not None
        for frame_index in range(frame_count):
            # Leave the final 0.5 s on the complete diagram for explanation.
            progress = min(1.0, frame_index / (frame_count - fps / 2))
            frame = _animation_frame(segment, number, total, progress)
            if frame_index == frame_count - 1:
                frame.save(poster)
            process.stdin.write(frame.tobytes())
        process.stdin.close()
        process.wait(timeout=120)
        stderr = process.stderr.read().decode("utf-8", errors="replace") if process.stderr else ""
    except Exception:
        process.kill()
        process.wait(timeout=10)
        raise
    if process.returncode != 0 or not output.is_file() or output.stat().st_size < 1000:
        raise PresentationError(f"动画片段合成失败：{stderr[-500:]}")


def _background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _color(INK)


def _textbox(slide, text: str, x: float, y: float, w: float, h: float,
             *, size: int, color: str = WHITE, bold: bool = False,
             align=PP_ALIGN.LEFT) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = _safe_text(text)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _color(color)


def _notes(slide, segment: LessonSegment, *, animation: bool,
           spoken_steps: list[str] | None = None) -> None:
    frame = slide.notes_slide.notes_text_frame
    if frame is not None:
        heading = "动画演示讲解" if animation else "教学图讲解"
        frame.text = (f"{heading}\n{_safe_text(segment.narration)}\n\n"
                      f"来源：PDF 第 {segment.evidence.page} 页\n"
                      f"原文摘录：{_safe_text(segment.evidence.quote)}")
        if segment.math_scene:
            import json
            frame.text += ("\n\n条件与约定：二维欧氏平面、列向量、标准输入/输出基；角度以度计。"
                           "投影为过原点直线上的正交投影；换基时依次使用平行、垂直方向。"
                           "一般矩阵的中间形变表示过渡，不宣称每个中间帧保留面积或长度。\n"
                           "教学参数："+json.dumps(segment.math_scene.parameters.model_dump(),ensure_ascii=False))
            frame.text += "\n\n数学推演分镜（教学示例，计算已核验）：\n" + "\n".join(
                f"{index+1}. {beat.action}：{beat.narration}"
                for index, beat in enumerate(segment.math_scene.beats))
        if segment.visual_scene:
            import json
            scene=segment.visual_scene
            frame.text += ("\n\n通用场景参数："+json.dumps(scene.parameters,ensure_ascii=False)+
                "\n示意与简化条件："+"；".join(scene.simplifications)+
                ("\n核验范围：原文摘录、关系引用与讲解状态；领域事实和教学解释需要复核。\n" if scene.diagram
                 else "\n核验范围：表达式、几何与声明的数值关系；领域事实和教学解释需要复核。\n")+
                "\n".join(f"{i+1}. {words}" for i, words in enumerate(
                    spoken_steps or [b.narration for b in scene.beats])))
            if scene.diagram:
                frame.text += '\n\n教学表达：'+scene.diagram.representation+'\n选择依据：'+scene.diagram.rationale
                frame.text += '\n节点与关系的原文依据：\n'+'\n'.join(
                    item.label+'：'+' '.join([item.source_quote,*getattr(item,'supporting_quotes',[])])
                    for item in [*scene.diagram.nodes,*scene.diagram.relations])


def _set_title(slide, segment: LessonSegment, number: int, total: int,
               *, animation: bool) -> None:
    _background(slide)
    _textbox(slide, segment.title, 0.66, 0.25, 10.8, 0.7,
             size=32, bold=True)
    label = ("点击播放数学推演 · 含配音" if animation else "SVG 数学摘要") if segment.math_scene else (
        ("点击播放教学过程 · 含配音" if animation else "SVG 场景摘要") if segment.visual_scene else
        ("点击播放动画" if animation else "SVG 教学图"))
    _textbox(slide, f"{number:02d} / {total:02d}  ·  {label}",
             0.69, 1.03, 11.8, 0.35, size=16, color=ACCENT, bold=True)
    _textbox(slide, f"来源：PDF 第 {segment.evidence.page} 页",
             0.69, 7.05, 5.8, 0.25, size=13, color=MUTED)


def _patch_svg_parts(original: Path, output: Path,
                     assets: dict[int, tuple[int, bytes]]) -> None:
    """Add native Office SVG references while retaining PNG fallback pictures."""
    with zipfile.ZipFile(original, "r") as source, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        overrides: dict[str, bytes] = {}
        content_root = etree.fromstring(source.read("[Content_Types].xml"))
        if not any(node.get("Extension") == "svg" for node in content_root):
            etree.SubElement(content_root, f"{{{CT_NS}}}Default",
                             Extension="svg", ContentType="image/svg+xml")
        overrides["[Content_Types].xml"] = etree.tostring(
            content_root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        for slide_number, entries in assets.items():
            # Accept legacy one-picture tuples and new multi-illustration pages.
            if isinstance(entries, tuple) and len(entries) == 2 and isinstance(entries[0], int):
                entries = [entries]
            slide_name = f"ppt/slides/slide{slide_number}.xml"
            rel_name = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
            slide = etree.fromstring(source.read(slide_name))
            rels = etree.fromstring(source.read(rel_name))
            for asset_number, (shape_id, svg_bytes) in enumerate(entries, 1):
                picture = slide.xpath(".//p:pic[p:nvPicPr/p:cNvPr[@id=$id]]",
                                      namespaces={"p": P_NS}, id=str(shape_id))
                if len(picture) != 1:
                    raise PresentationError("无法定位 SVG 教学图的图片占位符。")
                blips = picture[0].xpath("./p:blipFill/a:blip", namespaces={"p": P_NS, "a": A_NS})
                if len(blips) != 1:
                    raise PresentationError("SVG 教学图缺少 PNG 兼容图片。")
                used_ids = [int(match.group(1)) for rel in rels
                            if (match := re.fullmatch(r"rId(\d+)", rel.get("Id", "")))]
                r_id = f"rId{max(used_ids, default=0)+1}"
                media_name = f"teaching-{slide_number}-{asset_number}.svg"
                etree.SubElement(rels, f"{{{REL_NS}}}Relationship", Id=r_id,
                                 Type=REL_IMAGE, Target=f"../media/{media_name}")
                ext_list = blips[0].find(f"{{{A_NS}}}extLst")
                if ext_list is None:
                    ext_list = etree.SubElement(blips[0], f"{{{A_NS}}}extLst")
                ext = etree.SubElement(ext_list, f"{{{A_NS}}}ext", uri=SVG_EXT_URI)
                svg_blip = etree.SubElement(ext, f"{{{ASVG_NS}}}svgBlip", nsmap={"asvg": ASVG_NS})
                svg_blip.set(f"{{{R_NS}}}embed", r_id)
                overrides[f"ppt/media/{media_name}"] = svg_bytes
            overrides[slide_name] = etree.tostring(
                slide, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            overrides[rel_name] = etree.tostring(
                rels, xml_declaration=True, encoding="UTF-8", standalone=True
            )
        for item in source.infolist():
            target.writestr(item, overrides.pop(item.filename, source.read(item.filename)))
        for name, data in overrides.items():
            target.writestr(name, data)


def render_presentation(lesson: Lesson, output: Path) -> None:
    """Create a self-contained PPTX with one diagram and one animation per segment.

    The visual diagram is embedded as an SVG object plus a PNG compatibility
    fallback. Titles and footers are native editable text. Each animation is a
    self-contained MP4 media shape playable in slideshow mode (scene clips include audio). Narration
    and exact PDF excerpts are in speaker notes.
    """
    if lesson.ppt_template != 'classic' or lesson.deck_plan:
        from zhijiang.deck_renderer import render_template_presentation
        render_template_presentation(lesson, output)
        return
    if not lesson.segments:
        raise PresentationError("没有可制作成幻灯片的讲解片段。")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="book2course-pptx-", dir=output.parent) as temporary:
        work = Path(temporary)
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = _safe_text(lesson.title)
        presentation.core_properties.subject = "PDF 教学课件"
        presentation.core_properties.comments = "包含来源可追溯的 SVG 教学图和 MP4 动画。"
        blank = presentation.slide_layouts[6]
        cover = presentation.slides.add_slide(blank)
        _background(cover)
        bar = cover.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     Inches(0.75), Inches(1.45), Inches(0.09), Inches(3.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _color(ACCENT)
        bar.line.fill.background()
        _textbox(cover, "Book2Course · 教学课件", 1.03, 1.4, 8.8, 0.45,
                 size=19, color=ACCENT, bold=True)
        _textbox(cover, lesson.title, 1.03, 2.0, 11.4, 1.4,
                 size=44, bold=True)
        _textbox(cover, lesson.objective, 1.03, 3.7, 10.6, 1.15,
                 size=24, color=MUTED)
        _textbox(cover, f"{len(lesson.segments)} 个知识点 · SVG 图示与动画演示",
                 1.03, 6.65, 10.5, 0.35, size=17, color=ACCENT)
        cover_notes = cover.notes_slide.notes_text_frame
        if cover_notes is not None:
            cover_notes.text = f"课程目标：{_safe_text(lesson.objective)}"
            attribution = lesson.animation_report.get("source_attribution", {})
            if attribution:
                credit = f"原理来源：{attribution['author']} · {attribution['title']} · {attribution['license']}"
                _textbox(cover, credit, 1.03, 5.65, 11.4, 0.6, size=15, color=MUTED)
                cover_notes.text += f"\n{credit}\n{attribution['url']}\n{attribution['license_url']}\n本课数值例子、动画和中文解释为补充教学示例。"

        svg_assets: dict[int, tuple[int, bytes]] = {}
        spoken_by_segment = {
            asset.get("segment_index", index): [cue["text"] for cue in asset.get("timing", [])]
            for index, asset in enumerate(lesson.animation_report.get("assets", []))
        }
        total = len(lesson.segments)
        for index, segment in enumerate(lesson.segments, start=1):
            math_folder = output.parent / ("math" if segment.math_scene else "visual") / f"scene-{index:02d}"
            if segment.math_scene or segment.visual_scene:
                diagram_png = math_folder / "summary.png"
                svg_bytes = (math_folder / "summary.svg").read_bytes()
                diagram_height = 5.0
            else:
                diagram_png = work / f"diagram-{index}.png"
                _diagram_png(segment, diagram_png)
                svg_bytes = _diagram_svg(segment)
                diagram_height = 4.5
            diagram_slide = presentation.slides.add_slide(blank)
            _set_title(diagram_slide, segment, index, total, animation=False)
            with Image.open(diagram_png) as image:
                ratio=image.width/image.height
            picture_width=min(12.0,diagram_height*ratio)
            picture_height=picture_width/ratio
            picture = diagram_slide.shapes.add_picture(
                str(diagram_png), Inches((13.333-picture_width)/2),
                Inches(1.55+(diagram_height-picture_height)/2),
                width=Inches(picture_width), height=Inches(picture_height)
            )
            svg_assets[len(presentation.slides)] = (picture.shape_id, svg_bytes)
            _notes(diagram_slide, segment, animation=False,
                   spoken_steps=spoken_by_segment.get(index - 1))

            if segment.math_scene or segment.visual_scene:
                clip = math_folder / "clip.mp4"
                poster = math_folder / "poster.png"
            else:
                clip = work / f"animation-{index}.mp4"
                poster = work / f"animation-{index}.png"
                _render_animation(segment, index, total, clip, poster)
            animation_slide = presentation.slides.add_slide(blank)
            _set_title(animation_slide, segment, index, total, animation=True)
            animation_slide.shapes.add_movie(
                str(clip), Inches(1.9915), Inches(1.48), Inches(9.35), Inches(5.2594),
                poster_frame_image=str(poster), mime_type="video/mp4"
            )
            _notes(animation_slide, segment, animation=True,
                   spoken_steps=spoken_by_segment.get(index - 1))

        original = work / "presentation-base.pptx"
        presentation.save(original)
        patched = work / "presentation-final.pptx"
        _patch_svg_parts(original, patched, svg_assets)
        os.replace(patched, output)
