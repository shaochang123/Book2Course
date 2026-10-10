"""One composition for editable PPT pages, reveal frames and narrated clips."""
from __future__ import annotations
import html
import json
import os
from pathlib import Path
import tempfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from zhijiang.deck_planning import DeckPageChoice, validate_choices, _validate_hook, visual_details
from zhijiang.deck_renderer import Page, fit_picture, PAGE_W, PAGE_H
from zhijiang.presentation import PresentationError, _font, _notes, _patch_svg_parts, _render_animation
from zhijiang.presentation_templates import get_template
from zhijiang.visual_assets import AssetCatalog
from zhijiang.page_media import narrated_pages, scene_on_page, media_fingerprint, clip_duration


def story_illustration_bindings(visual, refs, narration):
    from zhijiang.caption_grounding import speech_sentences
    sentences=speech_sentences(narration);bindings={}
    for ref in refs:
        if not ref.anchor_text:continue
        direct=[i for i,item in enumerate(visual['items'])
                if ref.anchor_text in item['label']+' '+item.get('caption','')]
        # A graph node denotes its named entity. A different entity mentioned
        # in the same support sentence must not become this node's portrait.
        candidates=direct
        if not candidates and visual['representation'] not in {'process','relationship'}:
            candidates=[i for i,item in enumerate(visual['items'])
                if ref.anchor_text in sentences[item['sentence_id']-1]]
        for i in candidates:
            if i not in bindings:
                bindings[i]=ref;break
    return bindings


def wrap_lines(text, width, size=24):
    draw = ImageDraw.Draw(Image.new('RGB', (1,1)))
    font = _font(round(size*96/72))
    result, line = [], ''
    for char in text:
        if char == '\n':
            result.append(line); line = ''; continue
        if line and draw.textlength(line+char, font=font) > width*96:
            result.append(line); line = ''
        line += char
    if line:
        result.append(line)
    return result


def text_pages(captions, width, height, size=24, gap=.33):
    """Fit by pagination, never by shrinking or ellipsizing teaching conditions."""
    line_height = round(size*96/72)*1.25/96
    pages, page, used = [], [], 0
    for caption in captions:
        lines = wrap_lines(caption, width, size)
        full_capacity = max(1, int((height-.12)/line_height))
        remaining = max(0, int((height-used-.12)/line_height))
        # A sentence that fits a fresh page must stay whole. Filling the last
        # line of a busy page used to leave a one-character continuation.
        if used and len(lines) > remaining and len(lines) <= full_capacity:
            pages.append(page); page = []; used = 0
        while lines:
            capacity = max(0,int((height-used-.12)/line_height))
            if used and capacity < 1:
                pages.append(page); page = []; used = 0
                capacity = max(1,int((height-.12)/line_height))
            take = min(len(lines), max(1,capacity))
            page.append('\n'.join(lines[:take])); lines = lines[take:]; used += take*line_height+gap
            if lines:
                pages.append(page); page = []; used = 0
    if page:
        pages.append(page)
    return pages or [[]]


def concise_group_pages(texts, width, height):
    """Always keep the verified key sentence; speech details are optional.

    Do not make an extra slide just to display a fragment of the narration.
    Full narration and evidence remain in every page's speaker notes.
    """
    selected = [texts[0]]
    if len(text_pages(selected, width, height, 24, gap=.24)) > 1:
        return text_pages(selected, width, height, 24, gap=.24)
    for detail in texts[1:]:
        proposal = [*selected, detail]
        if len(text_pages(proposal, width, height, 24, gap=.24)) == 1:
            selected = proposal
    return text_pages(selected, width, height, 24, gap=.24)


def header(page, title):
    left = 1.3 if page.template.id == 'editorial' else .8
    width = PAGE_W-left-.75
    size = page.template.title_size
    lines = wrap_lines(title, width, size)
    height = len(lines)*round(size*96/72)*1.25/96+.12
    page.text(title, (left,.38,width,height), size=size, min_size=size, bold=True)
    return max(1.6, .38+height+.2)


def card_svg(captions, template, width, height):
    """A plain evidence-grounded prompt card, without invented arrows or labels."""
    parts = []
    y = 24
    for caption in captions:
        lines = wrap_lines(caption, (width-56)/96, 24)
        h = len(lines)*40+24
        parts.append(f'<rect x="8" y="{y}" width="{width-16}" height="{h}" rx="18" fill="#{template.panel}"/>')
        for i,line in enumerate(lines):
            parts.append(f'<text x="28" y="{y+40+i*40}" fill="#{template.ink}" font-size="32" font-family="Microsoft YaHei">{html.escape(line)}</text>')
        y += h+18
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">'
            '<title>讲解要点</title><desc>来自本段已确认的画面短句，无新增关系。</desc>'+''.join(parts)+'</svg>').encode()


def illustration_regions(layout, left, body_y, body_h, count):
    """Knowledge takes most of the width; library pictures are small accents.

    Main means largest among the selected accents, not the main slide content.
    These fixed regions are also used by reveal frames and the complete video.
    """
    x = PAGE_W-2.85
    if layout == 'collage':
        size = min(1.45, (body_h-.4)/max(count, 1))
        return [(x+.3, body_y+.2+i*(size+.18), size, size) for i in range(count)]
    primary = min(2.15, body_h-1.4 if count>1 else body_h)
    center_y = body_y+max(.15, (body_h-primary)/2-.35)
    boxes = [(x, center_y, primary, primary)]
    boxes.extend((x+i*1.05, center_y+primary+.15,.85,.85) for i in range(count-1))
    return boxes


def semantic_panels(choice, left, y, height):
    """Content groups determine where pictures belong, without topic rules."""
    width=PAGE_W-left-.75; count=len(choice.visual_groups); gap=.32
    if choice.composition=='annotated' and count>1:
        a=width*.48
        return [(left,y+.15,a,height-.25),
                *[(left+a+gap,y+.15+i*(height-.1)/(count-1),width-a-gap,
                   (height-.1)/(count-1)-.22) for i in range(count-1)]]
    if choice.composition=='statement' and count>1:
        return [(left,y+.15,width,height*.46),
                *[(left+i*(width+gap)/(count-1),y+height*.52,(width-gap*(count-2))/(count-1),height*.45)
                   for i in range(count-1)]]
    return [(left+i*(width+gap)/count,y+.3,(width-gap*(count-1))/count,height-.55)
            for i in range(count)]


def asset_caption(asset):
    # Identify what the picture actually depicts. A source anchor may name a
    # related stage/part; it must not rename the drawing itself.
    return (asset.title.removeprefix('手绘').removeprefix('历史墨线')
            .removesuffix('插画').removesuffix('符号'))


def render_shared_deck(lesson, output, *, audio_files=None):
    import resvg_py
    output = Path(output); output.parent.mkdir(parents=True, exist_ok=True)
    template = get_template(lesson.ppt_template); plan = lesson.deck_plan
    story=plan.get('visual_story',{})
    navigation=story.get('navigation')
    if navigation:
        from zhijiang.slide_story import DeckNavigation, validate_navigation
        validate_navigation(DeckNavigation.model_validate(navigation),lesson)
    if plan.get('template_id') != template.id:
        raise PresentationError('PPT 模板与版式规划不一致。')
    choices = [DeckPageChoice.model_validate(item) for item in plan.get('pages', [])]
    validate_choices(choices, lesson, template, range(1,len(lesson.segments)+1))
    _validate_hook(plan.get('hook_segment_id'), lesson)
    if [p.segment_id for p in choices] != list(range(1,len(lesson.segments)+1)):
        raise PresentationError('PPT 页面必须保留已核验课程的顺序。')
    catalog = AssetCatalog()
    for choice in choices:
        for ref in choice.asset_refs:
            catalog.path(ref.asset_id)
    root = output.parent
    previews = root/'presentation-preview'; media = root/'presentation-media'
    previews.mkdir(exist_ok=True); media.mkdir(exist_ok=True)
    if audio_files is None:
        audio_files = [root/('math' if s.math_scene else 'visual')/f'scene-{i:02d}'/'narration.wav'
            if s.math_scene or s.visual_scene else root/f'audio-{i-1}.wav' for i,s in enumerate(lesson.segments,1)]
    if len(audio_files) != len(lesson.segments):
        raise PresentationError('配音数量与课程不一致。')
    deck = Presentation(); deck.slide_width = Inches(PAGE_W); deck.slide_height = Inches(PAGE_H)
    deck.core_properties.title = lesson.title
    deck.core_properties.subject = template.name
    deck.core_properties.comments = plan.get('planner','')+'；素材仅用于教学配图。'
    pages, segments, svgs = [], [], {}
    left = 1.3 if template.id == 'editorial' else .8
    def new(kind, segment_id=None):
        page = Page(deck, template, previews, len(deck.slides)+1)
        page.export_svg = True
        pages.append({'page':page.number,'kind':kind, 'segment_id':segment_id})
        return page
    def picture(page, path, box, svg=None):
        shape = page.picture(path,box)
        if svg and getattr(page, 'export_svg', False):
            svgs.setdefault(page.number, []).append((shape.shape_id, svg))
        return shape
    def asset(page, ref, box):
        path = catalog.path(ref.asset_id)
        raster = catalog.rasterize(ref.asset_id, root/'visual-cache')
        picture(page,raster,box,path.read_bytes() if path.suffix == '.svg' else None)
        if getattr(page, 'export_svg', False):
            page.asset_regions = getattr(page, 'asset_regions', []) + [{'asset_id':ref.asset_id,'box':box}]
    def write_captions(page, captions, x,y,w, reveal, size=24):
        for i,text in enumerate(captions):
            h = len(text.splitlines())*round(size*96/72)*1.25/96+.1
            if i < reveal:
                page.rect((x,y,.055,h),template.accent)
                page.text(text,(x+.2,y,w-.2,h),size=size,min_size=size,bold=(i==0))
            y += h+.23
    def notes(page, segment, choice):
        _notes(page.slide, segment, animation=False)
        refs = choice.asset_refs
        if refs:
            page.slide.notes_slide.notes_text_frame.text += '\n教学配图（非来源证据）：'+json.dumps(
                [catalog.assets[r.asset_id].public() for r in refs],ensure_ascii=False)
    cover = new('cover')
    cover.rect((left,1.12,10.6,.09),template.accent)
    title_lines = wrap_lines(lesson.title,10.6,44)
    th = len(title_lines)*59*1.25/96+.15
    cover.text(lesson.title,(left,1.45,10.6,th),size=44,min_size=44,bold=True)
    y = max(3.8,1.45+th+.3)
    cover_text=' · '.join(g['label'] for g in navigation['groups']) if navigation else lesson.objective
    objective_parts = text_pages([cover_text],10.6,6.95-y,24)
    for text in objective_parts[0]:
        h = len(text.splitlines())*40/96+.1
        cover.text(text,(left,y,10.6,h),size=24); y += h
    cover.slide.notes_slide.notes_text_frame.text = lesson.objective+'\n'+lesson.notice+'\n'+json.dumps(
        lesson.animation_report.get('source_attribution',{}),ensure_ascii=False)
    if navigation:
        refs=next((c.asset_refs for c in choices if c.asset_refs),[])
        if refs:asset(cover,refs[0],(left,5.1,1.05,1.05))
        cover.text('从问题出发，理解概念与实例',(left+1.35,5.35,8.9,.9),size=26,min_size=26)
    cover.save_preview()
    for group in objective_parts[1:]:
        page=new('objective'); body_y=header(page,'学习目标')
        write_captions(page,group,left,body_y,10.7,len(group)); page.save_preview()
        page.slide.notes_slide.notes_text_frame.text=lesson.objective
    hook = lesson.segments[plan['hook_segment_id']-1]
    question_text = hook.visual_scene.question if hook.visual_scene else '如何理解“'+hook.title+'”？'
    for part in text_pages([question_text],10.3,4.7,32):
        question = new('question'); body_y = header(question,'先想一个问题')
        question.rect((left,body_y,10.8,4.9),template.panel,rounded=True)
        question.text('\n'.join(part),(left+.3,body_y+.35,10.3,4.4),size=32,min_size=32,bold=True)
        _notes(question.slide,hook,animation=False); question.save_preview()
    if navigation:
        route_labels=[g['label'] for g in navigation['groups']]
    else:
        # Legacy exports also avoid a last roadmap page with one dangling title.
        size=max(1,(len(lesson.segments)+2)//3)
        route_labels=[lesson.segments[i].title for i in range(0,len(lesson.segments),size)]
    for group in text_pages(route_labels,10.5,5.0):
        page=new('roadmap'); y=header(page,'这节课的路线')
        for i,label in enumerate(group):
            row_y=y+.2+i*1.25
            page.rect((left,row_y,.8,.75),template.accent,rounded=True)
            page.text(str(i+1),(left+.23,row_y+.11,.45,.55),size=26,color=template.background,bold=True)
            page.text(label,(left+1.1,row_y+.08,9.2,.85),size=30,min_size=30,bold=True)
        page.slide.notes_slide.notes_text_frame.text=json.dumps(navigation or {'groups':route_labels},ensure_ascii=False)
        page.save_preview()
    for choice in choices:
        index=choice.segment_id; segment=lesson.segments[index-1]
        computed = bool(segment.math_scene or segment.visual_scene or index in lesson.animation_report.get('prepared_scenes',[]))
        captions = [segment.bullets[i-1] for i in choice.bullet_ids]
        probe_height = len(wrap_lines(segment.title,PAGE_W-left-.75,template.title_size))*round(template.title_size*96/72)*1.25/96+.12
        body_y=max(1.6,.38+probe_height+.2); body_h=6.8-body_y
        refs=choice.asset_refs
        story_visual=story.get('visuals',{}).get(str(index),{}).get('visual') if not computed else None
        if story_visual:
            from zhijiang.slide_story import SlideVisualDraft, validate_visual
            validate_visual(SlideVisualDraft.model_validate(story_visual),segment)
        semantic=bool(choice.visual_groups) and not computed and not story_visual
        specs=[]
        if story_visual:
            groups=[[]]
        elif semantic:
            details={item['id']:item['text'] for item in visual_details(segment)}
            for group_id,(group,box) in enumerate(zip(choice.visual_groups,semantic_panels(choice,left,body_y,body_h)),1):
                bound=[ref for ref in refs if ref.group_id==group_id]
                texts=list(dict.fromkeys([segment.bullets[group.bullet_id-1],*[details[i] for i in group.detail_ids]]))
                aliases=list(dict.fromkeys(asset_caption(catalog.assets[ref.asset_id]) for ref in bound))
                label=' · '.join(aliases)
                label_lines=wrap_lines(label,box[2]-.6,24) if label else []
                art_height=(1.0+.52*len(label_lines)+.23) if bound else 0.
                text_height=box[3]-.5-art_height
                if text_height<.65:
                    # Use columns rather than compressing narrow annotation rows.
                    choice.composition='cards'
                    specs=[]
                    break
                parts=concise_group_pages(texts,box[2]-.55,text_height)
                specs.append({'box':box,'refs':bound,'label':label,'art_height':art_height,'parts':parts})
            if not specs:
                for group_id,(group,box) in enumerate(zip(choice.visual_groups,semantic_panels(choice,left,body_y,body_h)),1):
                    bound=[ref for ref in refs if ref.group_id==group_id]
                    texts=list(dict.fromkeys([segment.bullets[group.bullet_id-1],*[details[i] for i in group.detail_ids]]))
                    label=' · '.join(dict.fromkeys(asset_caption(catalog.assets[ref.asset_id]) for ref in bound))
                    art_height=(1.+.52*len(wrap_lines(label,box[2]-.6,24))+.23) if bound else 0.
                    specs.append({'box':box,'refs':bound,'label':label,'art_height':art_height,
                                  'parts':concise_group_pages(texts,box[2]-.55,box[3]-.5-art_height)})
            groups=[[] for _ in range(max(len(spec['parts']) for spec in specs))]
        elif computed:
            groups=[captions]; text_width=10.8; text_height=body_h
        elif refs:
            text_width=PAGE_W-left-3.45; text_height=body_h-.35
            groups=text_pages(captions,text_width-.3,text_height,size=28)
        else:
            original=template.boxes[choice.layout]
            text_width=original[2]-.65; text_height=min(body_h,original[3])-.5
            groups=text_pages(captions,text_width,text_height,gap=.45)
        states=[]; body_numbers=[]; viewport=None
        for group_number, group in enumerate(groups):
            active_specs=[]
            if semantic:
                active=[(i,spec) for i,spec in enumerate(specs) if group_number<len(spec['parts'])]
                partial=choice.model_copy(update={'visual_groups':[choice.visual_groups[i] for i,_ in active]})
                for (_,spec),box in zip(active,semantic_panels(partial,left,body_y,body_h)):
                    active_specs.append({**spec,'box':box})
            # Use exactly the same commands for native elements and raster states.
            def compose(page, reveal):
                nonlocal viewport
                header(page,segment.title)
                if story_visual:
                    from zhijiang.story_svg import visual_svg
                    bindings=story_illustration_bindings(story_visual,refs,segment.narration)
                    # Pictures sit beside their own object/example. They do not
                    # float in the title bar or replace the main teaching graph.
                    record_art=story_visual['representation']=='table' and bool(bindings)
                    width=PAGE_W-left-.75-(1.4 if record_art else 0)
                    box=(left,body_y,width,body_h-.1)
                    svg,node_boxes=visual_svg(story_visual,template,round(width*96),round(box[3]*96),
                                             reveal=reveal,art_slots=bindings)
                    path=media/f'story-{index}-{reveal}.png'
                    path.write_bytes(resvg_py.svg_to_bytes(svg_string=svg.decode(),font_family='Microsoft YaHei'))
                    picture(page,path,box,svg)
                    for art_index,(ni,ref) in enumerate(bindings.items()):
                        if ni>=reveal:continue
                        x,y,w,h=node_boxes[ni]
                        art_box=(left+width+.12,body_y+.4+art_index*1.2,1.05,1.05) if record_art else (
                            left+(x+w-80)/96,body_y+(y+16)/96,.65,.65)
                        asset(page,ref,art_box)
                elif semantic:
                    for gi,spec in enumerate(active_specs):
                        x,y,w,h=spec['box']
                        # Images and their original object names are within the
                        # same content group, rather than floating page furniture.
                        if gi>=reveal:
                            continue
                        # Comparisons benefit from a paired boundary; ordinary
                        # explanations keep open space rather than dashboard cards.
                        if choice.composition=='comparison' and len(active_specs)==2:
                            page.rect((x,y,w,h),template.panel,rounded=True)
                        page.rect((x+.22,y+.15,min(1.2,w-.44),.035),template.accent)
                        if spec['refs']:
                            for ai,ref in enumerate(spec['refs']):
                                asset(page,ref,(x+.2+ai*1.02,y+.28,.95,.95))
                            label_y=y+1.27
                            label_h=len(wrap_lines(spec['label'],w-.6,24))*.52+.14
                            page.text(spec['label'],(x+.25,label_y,w-.55,label_h),size=24,min_size=24)
                        ty=y+.26+spec['art_height']
                        for ti,text in enumerate(spec['parts'][group_number]):
                            th=len(text.splitlines())*40*1.25/96+.12
                            page.text(text,(x+.25,ty,w-.55,th),size=24,min_size=24,bold=ti==0)
                            ty+=th+.12
                elif computed:
                    folder=root/('math' if segment.math_scene else 'visual')/f'scene-{index:02d}'
                    image=folder/'summary.png'; svg=(folder/'summary.svg').read_bytes()
                    viewport=(left,body_y,PAGE_W-left-.6,max(.8,body_h-.15))
                    picture(page,image,viewport,svg)
                    for i,ref in enumerate(refs[:2]):
                        asset(page,ref,(left+i*.9,6.82,.7,.65))
                elif refs:
                    ordered=sorted(refs,key=lambda ref: ref.usage!='main')
                    # Only the content panel is prominent. Images occupy a
                    # narrow right margin and never replace an explanatory graph.
                    if choice.layout in {'evidence','split','illustrated'}:
                        first_h=len(group[0].splitlines())*47*1.25/96+.6
                        page.rect((left-.08,body_y+.05,text_width+.15,first_h),template.panel,rounded=True)
                    boxes=illustration_regions(choice.layout,left,body_y,body_h,len(ordered))
                    for i,(ref,box) in enumerate(zip(ordered,boxes)):
                        if i < reveal:
                            asset(page,ref,box)
                    write_captions(page,group,left+.08,body_y+.28,text_width-.1,reveal,size=28)
                else:
                    box=template.boxes[choice.layout]
                    viewport_box=(box[0],body_y,box[2],min(body_h,box[3]))
                    width,height=round(viewport_box[2]*96),round(viewport_box[3]*96)
                    svg=card_svg(group[:reveal],template,width,height)
                    path=media/f'card-{index}-{group_number}-{reveal}.png'
                    path.write_bytes(resvg_py.svg_to_bytes(svg_string=svg.decode(),font_family='Microsoft YaHei'))
                    picture(page,path,viewport_box,svg)
            reveal_count=len(story_visual['items']) if story_visual else len(active_specs) if semantic else max(len(group),len(refs),1)
            page=new('content',index); compose(page,reveal_count)
            notes(page,segment,choice); page.save_preview(); body_numbers.append(page.number)
            pages[-1].update(layout=choice.layout,asset_ids=[ref.asset_id for ref in refs],part=group_number+1,
                             composition=choice.composition,illustration_regions=getattr(page,'asset_regions',[]))
            if story_visual:
                pages[-1].update(representation=story_visual['representation'],
                    displayed_characters=sum(len(item['label'])+len(item['caption']) for item in story_visual['items'])+
                        len(story_visual['context'])+sum(len(e['label']) for e in story_visual['links']),
                    relation_count=len(story_visual['links']))
            if computed:
                states.append(previews/f'page-{page.number:03d}.png')
            else:
                for reveal in range(1,reveal_count+1):
                    # A temporary native page guarantees identical text placement.
                    temp_deck=Presentation()
                    state=Page(temp_deck,template,previews,page.number)
                    compose(state,reveal)
                    state_path=previews/f'page-{page.number:03d}-reveal-{reveal}.png'
                    state.preview.save(state_path); states.append(state_path)
        clip=media/f'segment-{index:03d}.mp4'; audio=Path(audio_files[index-1])
        duration=None
        if computed:
            source=root/('math' if segment.math_scene else 'visual')/f'scene-{index:02d}'/'clip.mp4'
            # Empty viewport background avoids a static summary behind letterboxing.
            back=Page(Presentation(),template,previews,0); header(back,segment.title)
            for i,ref in enumerate(refs[:2]):
                asset(back,ref,(left+i*.9,6.82,.7,.65))
            background=media/f'background-{index}.png'; back.preview.save(background)
            scene_on_page(background,source,viewport,clip)
            duration=clip_duration(clip)
        elif audio.is_file():
            duration=narrated_pages(states,audio,clip)
        elif choice.include_animation:
            # Legacy direct export can have no narration. New jobs always do.
            _render_animation(segment,index,len(lesson.segments),clip,media/f'poster-{index}.png')
        if clip.is_file():
            segments.append({'segment_id':index,'pages':body_numbers,'clip':clip.relative_to(root).as_posix(),
                             'duration':duration,'asset_ids':[r.asset_id for r in refs]})
        if choice.include_animation and clip.is_file():
            animation=new('animation',index)
            poster=previews/f'page-{body_numbers[-1]:03d}.png'
            animation.picture(poster,(0,0,PAGE_W,PAGE_H))
            pic=animation.slide.shapes[-1]; pic._element.getparent().remove(pic._element)
            animation.slide.shapes.add_movie(str(clip),Inches(0),Inches(0),Inches(PAGE_W),Inches(PAGE_H),
                poster_frame_image=str(poster),mime_type='video/mp4')
            _notes(animation.slide,segment,animation=True)
            animation.preview=Image.open(poster).convert('RGB'); animation.save_preview()
    if navigation:
        from zhijiang.story_svg import visual_svg
        recap=new('ending');y=header(recap,'带走这些核心结论')
        chosen=navigation['takeaway_ids'];items=[]
        for i in chosen:
            v=story.get('visuals',{}).get(str(i),{}).get('visual')
            if v:
                # The closing repeats reviewed assertions, never just a page
                # title such as "带走三句话". Conditions remain in the caption.
                item=dict(max(v['items'],key=lambda item:len(item['label'])+len(item['caption'])))
                if v.get('links'):
                    edge=v['links'][0]
                    item=dict(v['items'][edge['source']-1])
                    item['caption']=edge['label']+' → '+v['items'][edge['target']-1]['label']
                if v.get('context'):item['caption']+=(('；' if item['caption'] else '')+v['context'])
            else:
                s=lesson.segments[i-1]
                item={'label':s.title[:16],'caption':s.bullets[0]}
            items.append(item)
        if story.get('closing_points'):
            items=[{'label':p['label'],'caption':p['text']} for p in story['closing_points']]
        closing={'representation':'table' if story.get('closing_points') else
                 'comparison' if len(items)>1 else 'key_idea','items':items,'links':[],'context':''}
        svg,_=visual_svg(closing,template,round((PAGE_W-left-.75)*96),round((6.8-y)*96))
        path=media/'ending.png';path.write_bytes(resvg_py.svg_to_bytes(svg_string=svg.decode(),font_family='Microsoft YaHei'))
        picture(recap,path,(left,y,PAGE_W-left-.75,6.8-y),svg)
        recap.slide.notes_slide.notes_text_frame.text='\n\n'.join(
            lesson.segments[i-1].narration+'\n'+lesson.segments[i-1].evidence.model_dump_json() for i in chosen)
        recap.save_preview()
    else:
        recap=new('ending');y=header(recap,'回顾：你能解释了吗？')
        group=[s.title for s in lesson.segments[-3:]]
        write_captions(recap,group,left,y,10.7,len(group));recap.save_preview()
    bookends=[]
    if story:
        import wave
        for kind,duration in story.get('bookend_seconds',{}).items():
            num=next(p['page'] for p in pages if p['kind']==kind)
            silence=media/f'{kind}-silence.wav'
            with wave.open(str(silence),'wb') as audio:
                audio.setnchannels(1);audio.setsampwidth(2);audio.setframerate(24000)
                audio.writeframes(b'\x00\x00'*round(duration*24000))
            clip=media/f'{kind}.mp4'
            narrated_pages([previews/f'page-{num:03d}.png'],silence,clip)
            bookends.append({'kind':kind,'page':num,'duration':duration,'clip':clip.relative_to(root).as_posix(),
                             'audio':'intentional_silence_no_repeated_narration'})
    with tempfile.TemporaryDirectory(prefix='shared-ppt-',dir=root) as temp:
        raw=Path(temp)/'base.pptx'; final=Path(temp)/'final.pptx'
        deck.save(raw); _patch_svg_parts(raw,final,svgs); os.replace(final,output)
    manifest={'version':2,'template_id':template.id,'asset_catalog_fingerprint':catalog.fingerprint,
              'preview_type':'native-layout-review-not-powerpoint-render','pages':pages,'segments':segments,
              'bookends':bookends,
              'media_fingerprint':media_fingerprint(lesson,audio_files)}
    (previews/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return manifest
