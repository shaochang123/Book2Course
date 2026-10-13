"""Validate static vector courseware without embedded movie shapes."""

from __future__ import annotations

import zipfile
import subprocess
from pathlib import Path
from xml.etree import ElementTree as ET

import imageio_ffmpeg
from pptx import Presentation

from zhijiang.models import Evidence, Lesson, LessonSegment, Mode, VoiceMode
from zhijiang.presentation import render_presentation


def _lesson() -> Lesson:
    source = Evidence(page=2, quote="The probability lies between 0 and 1.")
    return Lesson(
        title="七年级概率导论",
        objective="用样本空间与概率范围解释简单随机实验。",
        segments=[
            LessonSegment(
                title="实验与样本空间",
                kind="concept",
                narration="先做一次随机实验。每一种可能结果都属于样本空间，结果本身只是一次输出。",
                bullets=["随机实验产生结果", "结果是一次实验的可能输出", "样本空间包含所有可能结果"],
                evidence=source,
            ),
            LessonSegment(
                title="概率的取值范围",
                kind="formula",
                narration="概率的数值从零到一。零表示不可能，一表示一定发生，中间表示不同的可能程度。",
                bullets=["概率在 0 到 1 之间", "0 表示不可能", "1 表示一定发生"],
                evidence=source,
            ),
            LessonSegment(
                title="计算顺序 <A & B>",
                kind="process",
                narration="先列出所有结果，然后数出目标事件中的结果，最后写出相应的概率比值。",
                bullets=["列出所有结果", "识别事件 <A>", "写出符合条件的结果数"],
                evidence=source,
            ),
        ],
        mode=Mode.DEMO,
        voice_mode=VoiceMode.SYSTEM,
        notice="演示",
    )


def test_pptx_embeds_static_svg_with_notes_and_bookends(tmp_path: Path) -> None:
    path = tmp_path / "lesson.pptx"
    render_presentation(_lesson(), path)

    deck = Presentation(path)
    assert len(deck.slides) == 5  # cover, one static diagram per point, closing
    for slide in (deck.slides[1],deck.slides[2],deck.slides[3]):
        picture=next(shape for shape in slide.shapes if shape.shape_type==13)
        pixels=picture.image.size
        assert abs(picture.width/picture.height-pixels[0]/pixels[1])<.0001
    assert "七年级概率导论" in " ".join(shape.text for shape in deck.slides[0].shapes
                                      if shape.has_text_frame)
    assert "计算顺序 <A & B>" in " ".join(shape.text for shape in deck.slides[3].shapes
                                              if shape.has_text_frame)

    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        names = archive.namelist()
        svg_parts = sorted(name for name in names if name.endswith(".svg"))
        movies = sorted(name for name in names if name.endswith(".mp4"))
        assert len(svg_parts) == 3
        assert movies == []
        svg_text = [archive.read(name).decode("utf-8") for name in svg_parts]
        for value in svg_text:
            ET.fromstring(value)
        assert any("样本空间包含所有可能结果" in value for value in svg_text)
        assert any("概率在 0 到 1 之间" in value for value in svg_text)
        assert not any("结果 1" in value or "可能的概率值" in value for value in svg_text)
        assert any("&lt;A &amp; B&gt;" in value for value in svg_text)

        drawing_ns = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
        relation_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        for slide_number in (2, 3, 4):
            xml = ET.fromstring(archive.read(f"ppt/slides/slide{slide_number}.xml"))
            svg_refs = xml.findall(f".//{{{drawing_ns}}}svgBlip")
            assert len(svg_refs) == 1
            rel_id = svg_refs[0].attrib[f"{{{relation_ns}}}embed"]
            rels = ET.fromstring(archive.read(
                f"ppt/slides/_rels/slide{slide_number}.xml.rels"
            ))
            matches = [rel for rel in rels if rel.attrib.get("Id") == rel_id]
            assert len(matches) == 1
            assert matches[0].attrib["Target"].endswith(".svg")
        notes = archive.read("ppt/notesSlides/notesSlide2.xml").decode("utf-8")
        assert "The probability lies between 0 and 1." in notes
        assert "样本空间" in notes
