import json
import subprocess
import wave
import zipfile
from xml.etree import ElementTree as ET
import pytest
import imageio_ffmpeg
from PIL import Image
from pptx import Presentation
from zhijiang.deck_planning import plan_deck, AssetReference
from zhijiang.deck_renderer import Page
from zhijiang.presentation import render_presentation, _patch_svg_parts
from zhijiang.presentation_templates import get_template
from zhijiang.video import render_video
from test_deck_planning import example_lesson

def audio_file(path,seconds=.6):
    with wave.open(str(path),'wb') as audio:
        audio.setnchannels(1);audio.setsampwidth(2);audio.setframerate(24000)
        audio.writeframes(b'\x00\x00'*round(seconds*24000))
    return path

def frame(path,time=0):
    result=subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(),'-v','error','-ss',str(time),
        '-i',str(path),'-frames:v','1','-f','rawvideo','-pix_fmt','rgb24','-'],capture_output=True,check=True)
    return Image.frombytes('RGB',(1280,720),result.stdout)

@pytest.mark.parametrize('template',['classic','paper','academic','editorial'])
def test_same_assets_in_ppt_and_video_and_no_small_system_labels(tmp_path,template):
    lesson=example_lesson(1);lesson.segments[0].title='西瓜的瓜瓤与瓜皮'
    lesson.segments[0].narration += '这里观察西瓜的瓜瓤和瓜皮，还可以观察瓜苗。'
    lesson.ppt_template=template;lesson.deck_plan=plan_deck(lesson,template)
    choice=lesson.deck_plan['pages'][0];choice['layout']='collage';choice['include_animation']=True
    choice['composition']='cards'
    choice['asset_refs']=[AssetReference(asset_id=key,anchor_text=anchor,usage=usage,group_id=i).model_dump()
        for i,(key,anchor,usage) in enumerate([('life-agriculture.watermelon','西瓜','main'),
             ('life-agriculture.watermelon-slice','瓜瓤','support'),('life-agriculture.melon-seedling','瓜苗','symbol')],1)]
    audio=audio_file(tmp_path/'audio-0.wav',1.2)
    render_presentation(lesson,tmp_path/'lesson.pptx');render_video(lesson,[audio],tmp_path/'lesson.mp4')
    deck=Presentation(tmp_path/'lesson.pptx')
    with zipfile.ZipFile(tmp_path/'lesson.pptx') as package:
        assert package.testzip() is None
        assert len([n for n in package.namelist() if n.endswith('.svg')])==3
        refs=ET.fromstring(package.read('ppt/slides/slide4.xml')).findall(
            './/{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip')
        assert len(refs)==3 and len({next(iter(r.attrib.values())) for r in refs})==3
    texts=' '.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    for label in ['AI 生成','需人工复核','SVG 可编辑','来源：PDF','BOOK2COURSE','点击播放']:
        assert label not in texts
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert all(run.font.size.pt>=24 for p in shape.text_frame.paragraphs for run in p.runs)
    manifest=json.loads((tmp_path/'presentation-preview/manifest.json').read_text(encoding='utf-8'))
    assert len(manifest['segments'])==1 and manifest['segments'][0]['duration']==pytest.approx(1.2)
    content=next(p for p in manifest['pages'] if p['kind']=='content')
    regions=content['illustration_regions']
    # All three decorative materials together remain below 8% of the canvas.
    assert len(regions)==3
    assert sum(r['box'][2]*r['box'][3] for r in regions)/(13.333333*7.5)<.08
    assert len({round(r['box'][0],1) for r in regions})==3  # each belongs to its content group
    # Last reveal and the actual video share pixel positions (allow H.264 loss).
    video=frame(tmp_path/'lesson.mp4',1.0)
    with Image.open(tmp_path/'presentation-preview/page-004.png') as page:
        import numpy as np
        assert abs(np.asarray(video).astype(float)-np.asarray(page).astype(float)).mean()<5
    notes=deck.slides[3].notes_slide.notes_text_frame.text
    assert lesson.segments[0].narration in notes and lesson.segments[0].evidence.quote in notes

def test_transparent_png_keeps_template_background(tmp_path):
    image=Image.new('RGBA',(100,100),(0,0,0,0));image.putpixel((50,50),(255,0,0,255))
    path=tmp_path/'image.png';image.save(path)
    page=Page(Presentation(),get_template('paper'),tmp_path,1)
    page.picture(path,(2,2,1,1))
    assert page.preview.getpixel((193,193))==(251,250,245)


def test_single_statement_does_not_get_an_empty_comparison_card(tmp_path):
    lesson=example_lesson(1);lesson.ppt_template='editorial'
    lesson.deck_plan=plan_deck(lesson,'editorial',use_illustrations=False)
    lesson.deck_plan['pages'][0].update(bullet_ids=[1],
        visual_groups=[{'bullet_id':1,'detail_ids':[]}],composition='comparison')
    audio=audio_file(tmp_path/'audio.wav',.2)
    from zhijiang.deck_renderer import render_template_presentation
    render_template_presentation(lesson,tmp_path/'lesson.pptx',audio_files=[audio])
    with Image.open(tmp_path/'presentation-preview/page-004.png') as page:
        color=get_template('editorial').background
        expected=tuple(int(color[i:i+2],16) for i in (0,2,4))
        assert page.getpixel((400,600))==expected


def test_key_sentences_do_not_leave_orphan_continuations():
    from zhijiang.rich_deck import text_pages, concise_group_pages
    first='前提条件需要完整说明'
    second='平均结论不能直接代替现实任务比较'
    pages=text_pages([first,second],4,1.4,24,gap=.24)
    assert len(pages)==2
    assert [''.join(page).replace('\n','') for page in pages]==[first,second]
    # Optional spoken detail does not create a mostly empty continuation.
    parts=concise_group_pages([first,second],4,1.1)
    assert len(parts)==1 and ''.join(parts[0]).replace('\n','')==first


def test_asset_caption_identifies_depicted_object_not_source_part():
    from zhijiang.rich_deck import asset_caption
    from zhijiang.visual_assets import AssetCatalog
    catalog=AssetCatalog()
    assert asset_caption(catalog.assets['medicine.lungs-ink'])=='肺部'
    assert asset_caption(catalog.assets['biology.butterfly-sketch'])=='蝴蝶成虫'
    assert asset_caption(catalog.assets['life-agriculture.watermelon-sketch'])=='西瓜切片'


@pytest.mark.parametrize('count',[1,2])
def test_last_reveal_is_held_for_the_entire_long_utterance(tmp_path,count):
    from zhijiang.page_media import narrated_pages
    audio=audio_file(tmp_path/'narration.wav',7.2)
    states=[]
    for i,color in enumerate(['red','blue'][:count]):
        path=tmp_path/f'{i}.png';Image.new('RGB',(1280,720),color).save(path);states.append(path)
    video=tmp_path/'clip.mp4';narrated_pages(states,audio,video)
    last=frame(video,6.9)
    assert last.getpixel((640,360))[0 if count==1 else 2]>240


@pytest.mark.parametrize('layout',['illustrated','split','evidence','captioned','collage'])
def test_illustrations_are_auxiliary_in_every_ordinary_layout(layout):
    from zhijiang.rich_deck import illustration_regions
    from zhijiang.deck_renderer import PAGE_W, PAGE_H
    for count in (1,2,3):
        regions=illustration_regions(layout,1.3,1.6,5.2,count)
        assert len(regions)==count
        assert sum(w*h for x,y,w,h in regions)<PAGE_W*PAGE_H*.08
        assert all(x>=PAGE_W-2.85 and y>=1.6 and x+w<=PAGE_W and y+h<7.5
                   for x,y,w,h in regions)

def test_long_conditions_paginate_without_ellipsis(tmp_path):
    lesson=example_lesson(1);lesson.ppt_template='paper'
    condition='只有当电压稳定、温度保持不变且测量仪器误差可接受时，结论才成立。'
    lesson.segments[0].bullets=[condition*15]
    lesson.deck_plan=plan_deck(lesson,'paper',use_illustrations=False)
    render_presentation(lesson,tmp_path/'lesson.pptx')
    manifest=json.loads((tmp_path/'presentation-preview/manifest.json').read_text(encoding='utf-8'))
    content=[p for p in manifest['pages'] if p['kind']=='content']
    assert len(content)>1
    with zipfile.ZipFile(tmp_path/'lesson.pptx') as package:
        svgs=[package.read(p).decode() for p in package.namelist() if p.endswith('.svg')]
        assert not any('…' in svg for svg in svgs)
        reconstructed=''.join(node.text or '' for svg in svgs for node in ET.fromstring(svg).findall(
            './/{http://www.w3.org/2000/svg}text'))
        assert reconstructed==condition*15

def test_existing_continuous_scene_is_retained_as_motion(tmp_path):
    from test_visual_scenes import scene_for
    lesson=example_lesson(1);lesson.ppt_template='editorial'
    lesson.segments[0].visual_scene=scene_for('数学','t',{'t':0},{'t':1})
    lesson.deck_plan=plan_deck(lesson,'editorial',use_illustrations=False)
    assert lesson.deck_plan['pages'][0]['layout']=='wide'
    folder=tmp_path/'visual/scene-01';folder.mkdir(parents=True)
    audio=audio_file(folder/'narration.wav',1.2)
    summary=Image.new('RGB',(640,360),'#FF0000');summary.save(folder/'summary.png')
    (folder/'summary.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 640 360"><rect width="640" height="360" fill="red"/></svg>')
    # A changing synthetic source tests preservation/compositing, not math truth.
    from zhijiang.page_media import narrated_pages
    blue=folder/'blue.png';Image.new('RGB',(640,360),'blue').save(blue)
    narrated_pages([folder/'summary.png',blue],audio,folder/'clip.mp4')
    render_presentation(lesson,tmp_path/'lesson.pptx');render_video(lesson,[audio],tmp_path/'lesson.mp4')
    first=frame(tmp_path/'lesson.mp4',.1);last=frame(tmp_path/'lesson.mp4',1)
    assert first.getpixel((640,400))[0]>200 and last.getpixel((640,400))[2]>200
    assert first.getpixel((1200,15))==last.getpixel((1200,15))

def test_create_and_retry_keep_illustration_choice(tmp_path,sample_pdf):
    from test_api import make_client
    client,store=make_client(tmp_path)
    with client:
        response=client.post('/api/jobs',files={'file':('sample.pdf',sample_pdf)},
            data={'rights_confirmed':'true','use_illustrations':'false','ppt_template':'paper'})
        job=response.json();assert job['model_settings']['use_illustrations'] is False
        store.fail(job['id'],'test failure')
        client.post('/api/jobs/'+job['id']+'/retry')
        assert store.get_options(job['id']).use_illustrations is False
        store.fail(job['id'],'test failure')
        client.post('/api/jobs/'+job['id']+'/retry',data={'use_illustrations':'true'})
        assert store.get_options(job['id']).use_illustrations is True
