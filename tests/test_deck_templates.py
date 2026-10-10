import json
import zipfile
from xml.etree import ElementTree as ET
import pytest
from PIL import Image
from pptx import Presentation
from zhijiang.deck_planning import plan_deck
from zhijiang.presentation import render_presentation
from zhijiang.presentation_templates import TEMPLATES
from test_deck_planning import example_lesson


def test_templates_have_distinct_page_geometry_not_only_palette():
    signatures = {tuple(sorted(template.boxes.items())) for template in TEMPLATES.values()}
    assert len(signatures) == len(TEMPLATES)


@pytest.mark.parametrize('template_id', ['paper', 'academic', 'editorial'])
def test_template_deck_is_self_contained_keeps_svg_media_and_notes(tmp_path, template_id):
    lesson = example_lesson(1)
    lesson.segments[0].kind = 'process'
    lesson.ppt_template = template_id
    lesson.deck_plan = plan_deck(lesson, template_id)
    path = tmp_path/f'{template_id}.pptx'
    render_presentation(lesson, path)
    deck = Presentation(path)
    assert len(deck.slides) == 6
    assert template_id in json.loads((tmp_path/'presentation-preview/manifest.json').read_text(encoding='utf-8'))['template_id']
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        svgs = [name for name in archive.namelist() if name.endswith('.svg')]
        movies = [name for name in archive.namelist() if name.endswith('.mp4')]
        assert len(svgs) == len(movies) == 1
        ET.fromstring(archive.read(svgs[0]))
        notes = archive.read('ppt/notesSlides/notesSlide4.xml').decode('utf-8')
        assert lesson.segments[0].evidence.quote in notes
        assert lesson.segments[0].narration in notes
    content = deck.slides[3]
    picture = next(shape for shape in content.shapes if shape.shape_type == 13)
    ratio = picture.image.size[0]/picture.image.size[1]
    assert abs(picture.width/picture.height-ratio) < 1e-5
    box = TEMPLATES[template_id].boxes[TEMPLATES[template_id].default_layout]
    from pptx.util import Inches
    assert picture.width <= Inches(box[2])+2
    assert picture.height <= Inches(box[3])+2
    for slide in deck.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left+shape.width <= deck.slide_width+3
            assert shape.top+shape.height <= deck.slide_height+3
    with Image.open(tmp_path/'presentation-preview/page-004.png') as preview:
        assert preview.size == (1280, 720)


def test_deck_template_mismatch_is_rejected_before_writing(tmp_path):
    lesson = example_lesson(1); lesson.ppt_template = 'paper'
    lesson.deck_plan = plan_deck(lesson, 'academic')
    from zhijiang.presentation import PresentationError
    with pytest.raises(PresentationError, match='不一致'):
        render_presentation(lesson, tmp_path/'invalid.pptx')


def test_definition_can_omit_extra_animation_page(tmp_path):
    lesson = example_lesson(1); lesson.ppt_template = 'paper'
    lesson.deck_plan = plan_deck(lesson, 'paper')
    assert not lesson.deck_plan['pages'][0]['include_animation']
    path = tmp_path/'definition.pptx'; render_presentation(lesson, path)
    assert len(Presentation(path).slides) == 5
    with zipfile.ZipFile(path) as archive:
        assert not any(name.endswith('.mp4') for name in archive.namelist())
