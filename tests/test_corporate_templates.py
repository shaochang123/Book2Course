import json
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET

from PIL import Image
from pptx import Presentation
import pytest

from zhijiang.deck_planning import plan_deck
from zhijiang.presentation import render_presentation
from zhijiang.presentation_templates import get_template,template_preview_svg,template_instruction
from test_deck_planning import example_lesson


@pytest.mark.parametrize('name',['newtouch_cloud','newtouch_blue','newtouch_white','newtouch_dark'])
def test_corporate_template_is_selectable_and_renders_same_background_in_native_ppt(name,tmp_path):
    template=get_template(name)
    ET.fromstring(template_preview_svg(name))
    assert template.name in template_instruction(name)
    lesson=example_lesson(1);lesson.ppt_template=name
    lesson.deck_plan=plan_deck(lesson,name,use_illustrations=False)
    output=tmp_path/'lesson.pptx';render_presentation(lesson,output)
    manifest=json.loads((tmp_path/'presentation-preview/manifest.json').read_text())
    assert manifest['template_id']==name
    deck=Presentation(output)
    assert len(deck.slides)>=5
    for slide in deck.slides:
        assert slide.shapes[0].shape_type==13  # Shared decorative backdrop.
        assert slide.shapes[0].image.size==(1280,720)
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=deck.slide_width+3
            assert shape.top+shape.height<=deck.slide_height+3
    with zipfile.ZipFile(output) as package:assert package.testzip() is None
    with Image.open(tmp_path/'presentation-preview/page-001.png') as preview:
        assert preview.size==(1280,720)
        assert len(set(preview.resize((80,45)).tobytes()))>100
    assert any(lesson.segments[0].narration in s.notes_slide.notes_text_frame.text for s in deck.slides)
