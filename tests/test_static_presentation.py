from zipfile import ZipFile
import json

from pptx import Presentation
import pytest

from zhijiang.agents import DemoAgents
from zhijiang.models import VoiceMode
from zhijiang.pdf import read_pdf
from zhijiang.presentation import render_presentation


@pytest.mark.parametrize('template', ['classic', 'academic'])
def test_ppt_is_static_with_notes_and_vectors_but_no_embedded_audio_or_video(tmp_path,sample_pdf,template):
    agents=DemoAgents();bundle=agents.extract_knowledge(read_pdf(sample_pdf,'source.pdf'))
    lesson=agents.script(bundle,agents.plan(bundle),VoiceMode.SYSTEM)
    lesson.ppt_template=template
    if template=='academic':
        from zhijiang.deck_planning import plan_deck
        lesson.deck_plan=plan_deck(lesson,template,use_illustrations=False)
    output=tmp_path/'lesson.pptx';render_presentation(lesson,output)
    with ZipFile(output) as archive:
        names=archive.namelist()
        assert any(name.endswith('.svg') for name in names)
        assert not any(name.endswith(('.mp4','.wav','.mp3','.m4a')) for name in names)
        assert not any(b'videoFile' in archive.read(name) or b'/relationships/video' in archive.read(name)
                       for name in names if name.endswith(('.xml','.rels')))
    deck=Presentation(output)
    notes='\n'.join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    assert all(segment.narration in notes for segment in lesson.segments)
    manifest=tmp_path/'presentation-preview/manifest.json'
    if manifest.is_file():assert not any(page['kind']=='animation' for page in json.loads(manifest.read_text(encoding='utf-8'))['pages'])
