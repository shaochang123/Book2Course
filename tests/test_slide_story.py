import json
import wave
from xml.etree import ElementTree as ET

import pytest
from PIL import Image
from pptx import Presentation

from zhijiang.agents import GenerationError
from zhijiang.slide_story import (SlideVisualDraft, VisualItem, VisualLink,
    DeckNavigation, RouteGroup, validate_navigation, validate_visual, plan_segment_visual,
    enrich_visual_story)
from zhijiang.deck_planning import plan_deck
from zhijiang.presentation_templates import get_template
from zhijiang.story_svg import visual_svg
from test_deck_planning import example_lesson


def example():
    lesson=example_lesson(1)
    lesson.segments[0].title='输入与结果'
    lesson.segments[0].narration='输入经过处理产生结果。处理条件保持不变时，比较才有意义。'
    return lesson


def diagram():
    return SlideVisualDraft(representation='process',items=[
        VisualItem(label='输入',caption='待处理的信息',sentence_id=1),
        VisualItem(label='结果',caption='处理后的输出',sentence_id=1)],
        links=[VisualLink(source=1,target=2,label='产生',sentence_id=1)],
        context='处理条件保持不变时，比较才有意义',context_sentence_id=2)


class Client:
    model='mock';base_url='localhost'
    def __init__(self,reject=False):self.calls=[];self.reject=reject
    def generate(self,schema,instruction,material):
        self.calls.append((schema.__name__,json.loads(material)))
        if schema.__name__=='SlideVisualDraft':return schema.model_validate(diagram().model_copy(update={'links':[]}).model_dump())
        if schema.__name__=='SlideLinkSelection':return schema.model_validate({
            'representation':'process','links':[v.model_dump() for v in diagram().links]})
        if schema.__name__=='SlideSourcePropositionsDraft':return schema.model_validate({'propositions':[]})
        if schema.__name__=='SlideVisualReview':
            fields={name:True for name in schema.model_fields if name!='issues'}
            fields['issues']=''
            if self.reject:fields.update(link_1=False,issues='方向错误')
            return schema.model_validate(fields)
        if schema.__name__=='DeckClosingReview':return schema.model_validate({'point_1':'knowledge_conclusion','issues':'输入和结果的关系是完整知识结论'})
        if schema.__name__=='DeckRouteReview':return schema.model_validate({
            **{name:'faithful_summary' for name in schema.model_fields if name!='issues'},'issues':'本组标题与输入和结果的关系一致'})
        return schema.model_validate({'cut_after':[],**{f'theme_{i}':'输入与结果' for i in range(1,5)},
            'closing_points':[{'segment_id':1,'sentence_id':1,'label':'输入'}]})


def test_compact_visual_keeps_script_and_uses_itemized_review(tmp_path):
    lesson=example();before=lesson.model_dump_json();client=Client()
    result=plan_segment_visual(lesson.segments[0],client,tmp_path/'visual.json')
    assert lesson.model_dump_json()==before
    assert result['review']['link_1'] and result['review']['context_complete']
    assert [c[0] for c in client.calls]==['SlideVisualDraft','SlideLinkSelection','SlideVisualReview']
    assert plan_segment_visual(lesson.segments[0],client,tmp_path/'visual.json')==result
    assert len(client.calls)==3


def test_reversed_or_unqualified_review_fails_closed(tmp_path):
    lesson=example();before=lesson.model_dump_json()
    with pytest.raises(GenerationError,match='方向错误'):
        plan_segment_visual(lesson.segments[0],Client(True),tmp_path/'visual.json')
    assert lesson.model_dump_json()==before and not (tmp_path/'visual.json').exists()
    assert len(json.loads((tmp_path/'visual.attempts.json').read_text(encoding='utf-8')))==3


def test_same_paragraph_does_not_authorize_a_link():
    s=example().segments[0];s.narration='输入属于资料。结果需要另外检查。'
    draft=diagram();draft.items[1].sentence_id=2
    with pytest.raises(ValueError,match='没有同时提及'):
        validate_visual(draft,s)


def test_invalid_nodes_or_arrowless_relationship_are_rejected():
    draft=diagram();draft.items[0].label='这些'
    with pytest.raises(ValueError,match='代词'):validate_visual(draft,example().segments[0])
    draft=diagram();draft.links=[]
    with pytest.raises(ValueError,match='需要有依据'):validate_visual(draft,example().segments[0])
    draft=diagram();draft.representation='comparison'
    with pytest.raises(ValueError,match='不能附加'):validate_visual(draft,example().segments[0])


def test_twenty_five_segments_have_one_coherent_grouped_route():
    lesson=example_lesson(25)
    nav=DeckNavigation(groups=[RouteGroup(label='理解概念',segment_ids=list(range(1,9))),
        RouteGroup(label='条件与例子',segment_ids=list(range(9,18))),
        RouteGroup(label='比较与总结',segment_ids=list(range(18,26)))],takeaway_ids=[3,13,20])
    validate_navigation(nav,lesson)
    nav.groups[-1].segment_ids=[25]
    with pytest.raises(ValueError,match='连续覆盖'):validate_navigation(nav,lesson)


def test_model_boundaries_expand_all_segments_once_without_listing_orphans():
    from types import SimpleNamespace
    from zhijiang.slide_story import navigation_from_boundaries
    selection=SimpleNamespace(cut_after=[17,8,8],theme_1='理解概念',theme_2='比较条件',theme_3='应用总结',
        theme_4='应用总结',closing_points=[SimpleNamespace(segment_id=2)])
    nav=navigation_from_boundaries(selection,example_lesson(25))
    assert [g.segment_ids for g in nav.groups]==[list(range(1,9)),list(range(9,18)),list(range(18,26))]
    selection.cut_after=[26]
    with pytest.raises(ValueError,match='真实片段之间'):navigation_from_boundaries(selection,example_lesson(25))
    selection.cut_after=[]
    with pytest.raises(ValueError,match='至少3个不同主题'):navigation_from_boundaries(selection,example_lesson(25))


def test_navigation_boundaries_cannot_be_omitted_and_silently_default_to_empty(tmp_path):
    from pydantic import ValidationError
    class BoundaryClient(Client):
        def generate(self,schema,instruction,material):
            if schema.__name__=='DeckNavigation':
                data={**{f'theme_{i}':'输入与结果' for i in range(1,5)},
                    'closing_points':[{'segment_id':1,'sentence_id':1,'label':'输入'}]}
                assert schema.model_fields['cut_after'].is_required()
                with pytest.raises(ValidationError):schema.model_validate(data)
            return super().generate(schema,instruction,material)
    lesson=example();lesson.deck_plan=plan_deck(lesson,'editorial')
    enrich_visual_story(lesson,BoundaryClient(),folder=tmp_path)


@pytest.mark.parametrize('view',['process','relationship','comparison','table','annotated','key_idea'])
@pytest.mark.parametrize('template',['classic','paper','academic','editorial'])
def test_diagram_views_render_transparent_editable_svg_without_invented_arrows(view,template):
    visual=diagram().model_dump();visual['representation']=view
    if view not in {'process','relationship'}:visual['links']=[]
    svg,boxes=visual_svg(visual,get_template(template),940,490)
    root=ET.fromstring(svg)
    assert root.find('{http://www.w3.org/2000/svg}title') is not None
    arrows=[p for p in root.findall('.//{http://www.w3.org/2000/svg}path') if p.get('marker-end')]
    assert len(arrows)==(1 if view in {'process','relationship'} else 0)
    assert all(float(t.get('font-size'))>=28 for t in root.findall('.//{http://www.w3.org/2000/svg}text'))
    import resvg_py,io
    im=Image.open(io.BytesIO(resvg_py.svg_to_bytes(svg_string=svg.decode())))
    assert im.mode=='RGBA' and im.getpixel((0,0))[3]==0
    assert len(boxes)==2


def test_model_failure_does_not_turn_into_a_demo_visual():
    class Fails:
        def generate(self,*a):raise GenerationError('服务不可用')
    lesson=example();lesson.deck_plan=plan_deck(lesson,'paper');before=lesson.model_dump_json()
    with pytest.raises(GenerationError,match='服务不可用'):enrich_visual_story(lesson,Fails())
    assert lesson.model_dump_json()==before


def test_independent_source_proposition_restores_graph_without_textbook_rules(monkeypatch,tmp_path):
    from zhijiang.slide_story import prefer_grounded_graph
    lesson=example();previous={'visual':diagram().model_copy(update={
        'representation':'annotated','links':[]}).model_dump(mode='json')}
    monkeypatch.setattr('zhijiang.teaching_graph.source_propositions',lambda *a,**kw:[{
        'subject':'输入','predicate':'产生','object':'结果','fact_ids':[1]}])
    result=prefer_grounded_graph(lesson.segments[0],previous,Client(),tmp_path/'graph.json')
    assert result['visual']['representation']=='relationship'
    assert [(e['source'],e['target'],e['label']) for e in result['visual']['links']]==[(1,2,'产生')]
    assert result['visual']['context']=='输入经过处理产生结果'
    assert result['independent_graph']['accepted']
    assert previous['visual']['links']==[]


def test_rejected_graph_alternative_keeps_reviewed_content_and_records_rejection(monkeypatch,tmp_path):
    from zhijiang.slide_story import prefer_grounded_graph
    previous={'visual':diagram().model_copy(update={'representation':'annotated','links':[]}).model_dump(mode='json')}
    monkeypatch.setattr('zhijiang.teaching_graph.source_propositions',lambda *a,**kw:[{
        'subject':'输入','predicate':'产生','object':'结果','fact_ids':[1]}])
    result=prefer_grounded_graph(example().segments[0],previous,Client(True),tmp_path/'graph.json')
    assert result['visual']==previous['visual']
    assert not result['independent_graph']['accepted']


def test_graph_layout_follows_real_direction_not_list_order():
    visual=diagram().model_dump();visual['links'][0].update(source=2,target=1)
    _,boxes=visual_svg(visual,get_template('paper'),940,490)
    assert boxes[1][0]<boxes[0][0]


def test_long_annotations_change_composition_instead_of_shrinking_text():
    visual=diagram().model_dump();visual.update(representation='annotated',links=[],context='')
    visual['items']=[{'label':'概念','caption':'这是必须完整保留的具体解释与条件以及对应的实际例子'} for _ in range(3)]
    _,boxes=visual_svg(visual,get_template('paper'),940,490)
    assert len({box[1] for box in boxes})==1


def test_comma_separated_claims_cannot_become_one_graph_edge(monkeypatch,tmp_path):
    from zhijiang.slide_story import prefer_grounded_graph
    lesson=example();lesson.segments[0].narration='聚类把相似对象分组，通常没有预先给出的结果标记。'
    previous={'visual':diagram().model_copy(update={'representation':'annotated','links':[]}).model_dump(mode='json')}
    monkeypatch.setattr('zhijiang.teaching_graph.source_propositions',lambda *a,**kw:[{
        'subject':'聚类','predicate':'把相似对象分组','object':'通常没有预先给出的结果标记','fact_ids':[1]}])
    result=prefer_grounded_graph(lesson.segments[0],previous,Client(),tmp_path/'graph.json')
    assert result['visual']==previous['visual'] and not result['independent_graph']['accepted']


def test_two_independent_facts_keep_comparison_instead_of_graph(monkeypatch,tmp_path):
    from zhijiang.slide_story import prefer_grounded_graph
    lesson=example();lesson.segments[0].narration='规则强调表达。模型需要参数。'
    previous={'visual':diagram().model_copy(update={'representation':'comparison','links':[]}).model_dump(mode='json')}
    monkeypatch.setattr('zhijiang.teaching_graph.source_propositions',lambda *a,**kw:[
        {'subject':'规则','predicate':'强调','object':'表达','fact_ids':[1]},
        {'subject':'模型','predicate':'需要','object':'参数','fact_ids':[2]}])
    result=prefer_grounded_graph(lesson.segments[0],previous,Client(),tmp_path/'graph.json')
    assert result['visual']==previous['visual'] and not result['independent_graph']['accepted']


def test_unresolved_plural_pronouns_cannot_label_objects():
    from zhijiang.teaching_graph import complete_graph_label
    assert not any(complete_graph_label(v) for v in ['它们','她们','自己','其'])


def test_visual_names_reject_loose_verbs_but_keep_nominal_operations():
    from zhijiang.slide_story import visual_object_name
    assert not visual_object_name('记住','学完后记住三条主线')
    assert not visual_object_name('建立','机器学习可以建立联系')
    assert not visual_object_name('演绎则','演绎则从原理推出结论')
    assert not visual_object_name('产生','从传感器信息产生方向控制')
    assert not visual_object_name('传感器信息产生','从传感器信息产生方向控制')
    assert not visual_object_name('理解','理解应用时先看输入和输出')
    assert not visual_object_name('问','理解应用时，先问输入是什么')
    assert not visual_object_name('氧气从肺泡','氧气从肺泡扩散进入血液')
    assert visual_object_name('回归','预测连续数值是回归')
    assert visual_object_name('判断','模型对新情况作出判断')
    assert visual_object_name('输入','输入经过处理产生结果')
    assert visual_object_name('输入','但是，输入经过处理产生结果')
    assert visual_object_name('敲声','敲声为浊响')


def test_corrupt_or_unapproved_cached_graph_is_not_reused(tmp_path):
    from zhijiang.slide_story import cached_alternative
    path=tmp_path/'graph.json';path.write_text('{bad',encoding='utf-8')
    assert cached_alternative(path,'key',example().segments[0],{},'independent_graph') is None
    d=diagram();review={name:True for name in __import__('zhijiang.slide_story',fromlist=['review_schema']).review_schema(d).model_fields if name!='issues'}
    review.update(link_1=False,issues='reversed')
    path.write_text(json.dumps({'fingerprint':'key','accepted':True,'visual':d.model_dump(),'review':review}),encoding='utf-8')
    assert cached_alternative(path,'key',example().segments[0],{},'independent_graph') is None


def test_cover_ending_and_shared_svg_are_in_ppt_and_full_video(tmp_path):
    from zhijiang.deck_renderer import render_template_presentation
    from zhijiang.video import render_video
    from zhijiang.page_media import clip_duration
    lesson=example();lesson.ppt_template='editorial';lesson.deck_plan=plan_deck(lesson,'editorial')
    enrich_visual_story(lesson,Client(),folder=tmp_path/'story')
    audio=tmp_path/'audio.wav'
    with wave.open(str(audio),'wb') as out:
        out.setnchannels(1);out.setsampwidth(2);out.setframerate(24000);out.writeframes(b'\0\0'*24000)
    result=render_template_presentation(lesson,tmp_path/'lesson.pptx',audio_files=[audio])
    render_video(lesson,[audio],tmp_path/'lesson.mp4')
    assert len([p for p in result['pages'] if p['kind']=='roadmap'])==1
    assert result['pages'][0]['kind']=='cover' and result['pages'][-1]['kind']=='ending'
    assert {p['kind'] for p in result['bookends']}=={'cover','ending'}
    assert clip_duration(tmp_path/'lesson.mp4')==pytest.approx(9,abs=.2)
    timeline=json.loads((tmp_path/'video-timeline.json').read_text(encoding='utf-8'))
    assert timeline['narration_uses_per_segment']==1
    deck=Presentation(tmp_path/'lesson.pptx')
    assert '输入经过处理产生结果' in deck.slides[-1].notes_slide.notes_text_frame.text
    from test_shared_deck import frame
    assert abs(__import__('numpy').asarray(frame(tmp_path/'lesson.mp4',.4)).astype(float)-
        __import__('numpy').asarray(Image.open(tmp_path/'presentation-preview/page-001.png')).astype(float)).mean()<5
    assert abs(__import__('numpy').asarray(frame(tmp_path/'lesson.mp4',7)).astype(float)-
        __import__('numpy').asarray(Image.open(tmp_path/'presentation-preview'/f'page-{len(deck.slides):03d}.png')).astype(float)).mean()<5


def test_all_continuous_scenes_still_get_source_bound_bookends(tmp_path):
    lesson=example();lesson.deck_plan=plan_deck(lesson,'paper')
    lesson.animation_report={'prepared_scenes':[1]}
    class SceneClient(Client):
        def generate(self,schema,instruction,material):
            if schema.__name__=='DeckNavigation':return schema.model_validate({
                'cut_after':[],**{f'theme_{i}':'输入与结果' for i in range(1,5)},
                'closing_points':[{'segment_id':1,'sentence_id':1,'label':'结果'}]})
            return super().generate(schema,instruction,material)
    story=enrich_visual_story(lesson,SceneClient(),folder=tmp_path)
    assert story['visuals']=={} and story['closing_points'][0]['text']=='输入经过处理产生结果'


def test_rejected_navigation_ending_does_not_publish_partial_plan(tmp_path):
    lesson=example();lesson.deck_plan=plan_deck(lesson,'paper');before=lesson.model_dump_json()
    class IntroEnding(Client):
        def generate(self,schema,instruction,material):
            if schema.__name__=='DeckClosingReview':return schema.model_validate({'point_1':'course_navigation','issues':'课程导航不是知识结论'})
            return super().generate(schema,instruction,material)
    with pytest.raises(GenerationError,match='课程导航不是知识结论'):
        enrich_visual_story(lesson,IntroEnding(),folder=tmp_path)
    assert lesson.model_dump_json()==before
    assert len(json.loads((tmp_path/'navigation-attempts.json').read_text(encoding='utf-8')))==1


def test_wrong_route_theme_is_replanned_and_cannot_publish(tmp_path):
    lesson=example();lesson.deck_plan=plan_deck(lesson,'paper');before=lesson.model_dump_json()
    class BadRoute(Client):
        def generate(self,schema,instruction,material):
            if schema.__name__=='DeckRouteReview':return schema.model_validate({
                'group_1':'unrelated','issues':'所选主题名不描述本组实际的输入和结果'})
            return super().generate(schema,instruction,material)
    with pytest.raises(GenerationError,match='路线主题复核拒绝'):enrich_visual_story(lesson,BadRoute(),folder=tmp_path)
    assert lesson.model_dump_json()==before
    attempts=json.loads((tmp_path/'navigation-attempts.json').read_text(encoding='utf-8'))
    assert len(attempts)==3 and all(a['route_review']['group_1']=='unrelated' for a in attempts)


def test_route_repair_keeps_reviewed_groups_and_boundaries(tmp_path):
    from typing import get_args
    lesson=example_lesson(5)
    for s in lesson.segments:s.narration=example().segments[0].narration;s.title='输入与结果'
    lesson.deck_plan=plan_deck(lesson,'paper')
    class FocusRoute(Client):
        routes=0;reviews=0
        def generate(self,schema,instruction,material):
            if schema.__name__=='DeckNavigation':
                self.routes+=1
                if self.routes==2:
                    assert get_args(schema.model_fields['theme_1'].annotation)==('输入处理',)
                    assert json.loads(material)['route_repair']['review']['group_2']=='misleading'
                return schema.model_validate({'cut_after':[2],'theme_1':'输入处理',
                    'theme_2':'结果记录' if self.routes==1 else '结果比较',
                    'theme_3':'结果记录','theme_4':'结果记录',
                    'closing_points':[{'segment_id':1,'sentence_id':1,'label':'输入'}]})
            if schema.__name__=='DeckRouteReview':
                self.reviews+=1
                return schema.model_validate({'group_1':'faithful_summary','group_2':'faithful_summary' if self.reviews>1 else 'misleading',
                    'issues':'第二组应概括比较的内容，第一组无需修改'})
            return super().generate(schema,instruction,material)
    client=FocusRoute();story=enrich_visual_story(lesson,client,folder=tmp_path)
    assert client.routes==2 and client.reviews==2
    assert [g['label'] for g in story['navigation']['groups']]==['输入处理','结果比较']
    assert story['navigation']['groups'][0]['segment_ids']==[1,2]


def test_repeated_labels_and_repeated_table_statements_are_not_teaching_examples():
    lesson=example();draft=diagram();draft.links=[];draft.representation='annotated'
    draft.items[1]=VisualItem(label='输入',caption='另外的解释',sentence_id=1)
    with pytest.raises(ValueError,match='短名重复'):validate_visual(draft,lesson.segments[0])
    draft=diagram();draft.links=[];draft.representation='table'
    for item in draft.items:item.caption='把同一整段类别列举重复在表格每一行'
    with pytest.raises(ValueError,match='类别列举'):validate_visual(draft,lesson.segments[0])
    # Equal short values are legitimate data, unlike repeated prose.
    for item in draft.items:item.caption='正常'
    validate_visual(draft,lesson.segments[0])


def test_screen_caption_cannot_end_at_a_cut_off_connective():
    draft=diagram();draft.items[0].caption='不能解决系统的安全与'
    with pytest.raises(ValueError,match='连接词处截断'):validate_visual(draft,example().segments[0])


def test_relation_candidates_preserve_active_clause_roles_instead_of_noun_cooccurrence():
    from zhijiang.slide_story import ordered_source_relation
    assert not ordered_source_relation('气体从容器进入管道','容器','气体','管道')
    assert not ordered_source_relation('气体从容器进入管道','容器','进入','气体')
    assert not ordered_source_relation('观察属于记录，分类需要证据','观察','需要','证据')
    assert ordered_source_relation('输入经过处理产生结果','输入','产生','结果')
    assert not ordered_source_relation('输入如何经过处理产生结果','输入','产生','结果')
    assert not ordered_source_relation('程序利用数据产生模型','程序','利用','模型')
    assert ordered_source_relation('程序利用数据产生模型','程序','利用数据产生','模型')


def test_branch_connectors_leave_tall_source_through_the_side_at_target_height():
    import re
    visual={'representation':'relationship','items':[
        {'label':'输入'},{'label':'结果甲'},{'label':'结果乙'}],
        'links':[{'source':1,'target':2,'label':'产生','kind':'directed'},
                 {'source':1,'target':3,'label':'产生','kind':'directed'}]}
    svg,boxes=visual_svg(visual,get_template('paper'),940,490)
    paths=[p for p in ET.fromstring(svg).findall('.//{http://www.w3.org/2000/svg}path') if p.get('marker-end')]
    for path,target in zip(paths,boxes[1:]):
        x1,y1,x2,y2=map(float,re.findall(r'-?\d+(?:\.\d+)?',path.get('d')))
        assert x1==pytest.approx(boxes[0][0]+boxes[0][2]) and x2==pytest.approx(target[0])
        assert y1==y2 and target[1]<y2<target[1]+target[3]


def test_rejected_object_is_removed_from_native_repair_contract(tmp_path):
    from pydantic import ValidationError
    class RepairClient(Client):
        attempts=0
        def generate(self,schema,instruction,material):
            if schema.__name__=='SlideVisualDraft':
                self.attempts+=1
                if self.attempts==2:
                    with pytest.raises(ValidationError):schema.model_validate(
                        diagram().model_copy(update={'links':[]}).model_dump())
                    return schema.model_validate({'representation':'key_idea',
                        'items':[{'label':'输入','caption':'待处理的信息','sentence_id':1}]})
            if schema.__name__=='SlideVisualReview' and self.attempts==1:
                return schema.model_validate({'item_1':True,'item_2':False,'link_1':True,
                    'context_complete':False,'representation_faithful':True,'issues':'结果的解释与可见条件不正确'})
            return super().generate(schema,instruction,material)
    client=RepairClient();result=plan_segment_visual(example().segments[0],client,tmp_path/'view.json')
    assert client.attempts==2 and len(result['visual']['items'])==1
    assert result['attempts'][0]['review']['item_2'] is False


def test_closing_cannot_depend_on_an_unnamed_previous_example():
    from zhijiang.slide_story import self_contained_closing
    for text in ['这个例子剩下三个规则','该实验说明现象','上述结果可以推广','另外两个不满足']:
        assert not self_contained_closing(text)
    assert self_contained_closing('模型对未见过的新样本作出判断')
    assert self_contained_closing('比较必须保持实验条件相同')


def test_repair_can_focus_on_approved_facts_but_must_review_the_reduced_page(tmp_path):
    class FocusClient(Client):
        reviews=0
        def generate(self,schema,instruction,material):
            if schema.__name__=='SlideVisualReview':
                self.reviews+=1
                if self.reviews==1:return schema.model_validate({'item_1':True,'item_2':False,'link_1':False,
                    'context_complete':True,'representation_faithful':True,'issues':'第二个对象的解释不完整'})
                assert len(json.loads(material)['items'])==1
            return super().generate(schema,instruction,material)
    client=FocusClient();result=plan_segment_visual(example().segments[0],client,tmp_path/'view.json')
    assert client.reviews==2 and result['visual']['links']==[]
    assert result['attempts'][-1]['repair_strategy']=='retain_reviewed_items'


def test_redundant_predicate_object_is_bound_without_inventing_relation(monkeypatch,tmp_path):
    from zhijiang.slide_story import prefer_grounded_graph
    previous={'visual':diagram().model_copy(update={'representation':'annotated','links':[]}).model_dump(mode='json')}
    monkeypatch.setattr('zhijiang.teaching_graph.source_propositions',lambda *a,**kw:[{
        'subject':'输入','predicate':'产生结果','object':'结果','fact_ids':[1]}])
    result=prefer_grounded_graph(example().segments[0],previous,Client(),tmp_path/'graph.json')
    assert result['visual']['links'][0]['label']=='产生'
    assert result['independent_graph']['selected'][0]['original_predicate']=='产生结果'


def test_graph_review_checks_named_nodes_with_incident_relations_not_missing_prose(monkeypatch,tmp_path):
    from zhijiang.slide_story import prefer_grounded_graph
    previous={'visual':diagram().model_copy(update={'representation':'annotated','links':[]}).model_dump(mode='json')}
    monkeypatch.setattr('zhijiang.teaching_graph.source_propositions',lambda *a,**kw:[{
        'subject':'输入','predicate':'产生','object':'结果','fact_ids':[1]}])
    class NodeClient(Client):
        def generate(self,schema,instruction,material):
            if schema.__name__=='SlideVisualReview':
                data=json.loads(material)
                assert all(n['role']=='object_identity' and n['incident_links'] for n in data['node_checks'])
                assert 'not a standalone factual assertion' in schema.model_fields['item_1'].description
                assert data['visual']['items'][0]['caption']==''
            return super().generate(schema,instruction,material)
    result=prefer_grounded_graph(example().segments[0],previous,NodeClient(),tmp_path/'graph.json')
    assert result['independent_graph']['accepted'] and result['visual']['links']


def test_pictures_prioritize_visible_object_not_first_shared_sentence():
    from types import SimpleNamespace
    from zhijiang.rich_deck import story_illustration_bindings
    visual={'representation':'annotated','items':[
        {'label':'规则','caption':'明确表示','sentence_id':1},
        {'label':'神经网络','caption':'模型','sentence_id':1}]}
    ref=SimpleNamespace(anchor_text='神经网络')
    assert story_illustration_bindings(visual,[ref],'规则和神经网络都是模型。')=={1:ref}
    visual={'representation':'relationship','items':[
        {'label':'电池','caption':'','sentence_id':1},{'label':'电压','caption':'','sentence_id':1}]}
    assert story_illustration_bindings(visual,[SimpleNamespace(anchor_text='灯泡')],
        '电池提供电压，灯泡发光。')=={}
    visual={'representation':'table','items':[
        {'label':'颜色','caption':'绿色','sentence_id':1}]}
    ref=SimpleNamespace(anchor_text='果实')
    assert story_illustration_bindings(visual,[ref],'果实记录：颜色绿色。')=={0:ref}


def test_closing_removes_discourse_but_preserves_negation_conditions_and_full_source(tmp_path):
    from zhijiang.slide_story import closing_display_text
    assert closing_display_text('第三，模型不能直接推广')=='模型不能直接推广'
    assert closing_display_text('但是，条件相同时才能比较')=='条件相同时才能比较'
    assert closing_display_text('因此，结论依赖前提')=='因此，结论依赖前提'
    lesson=example();lesson.segments[0].narration='但是，'+lesson.segments[0].narration
    story=enrich_visual_story(lesson,Client(),folder=tmp_path)
    assert story['closing_points'][0]['text']=='输入经过处理产生结果'
    assert story['closing_points'][0]['source_sentence']=='但是，输入经过处理产生结果'
