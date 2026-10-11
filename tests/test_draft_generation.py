import json

from zhijiang.draft_generation import build_draft_lesson, material_for_points
from zhijiang.models import Evidence, KnowledgeBundle, KnowledgePoint, PageText, SourceDocument, VoiceMode
from zhijiang.deck_planning import DeckPageChoice, validate_choices
from zhijiang.presentation_templates import get_template
from zhijiang.agents import validate_lesson


def inputs(count=5):
    # The definition crosses both a 300-character quotation and a page boundary.
    first='任务T是要完成的工作。性能P用来衡量任务完成的表现。'+('教材中的背景解释。'*40)
    second='经验E是学习时获得的资料。如果系统在任务T上的性能P随经验E提高，就称为学习。'
    document=SourceDocument(filename='公开原创讲义.pdf',pages=[PageText(page=1,text=first),PageText(page=2,text=second)])
    bundle=KnowledgeBundle(points=[KnowledgePoint(title=f'核心概念{i}',kind='concept',
        explanation='解释任务、性能和经验三者的关系，以及这个定义的条件。',
        evidence=Evidence(page=1,quote=first[:30])) for i in range(1,count+1)])
    return document,bundle


class DraftClient:
    model='test';base_url='http://localhost';semantic_thinking=False

    def __init__(self):self.calls=[]

    def generate(self,schema,instruction,material):
        data=json.loads(material);self.calls.append(data)
        if schema.__name__=='DraftBookends':
            return schema.model_validate({'cut_after':[], 'theme_labels':['学习的定义'],
                'closing_ids':[data['closing_candidates'][0]['id']]})
        assert schema.__name__=='TeachingDraftBatch'  # No completeness/approval veto call.
        assert data['source_pages'][0]['text'].endswith('教材中的背景解释。')
        assert '经验E' in data['source_pages'][1]['text']
        return schema.model_validate({'title':'理解学习的定义','objective':'理解任务、性能和经验的关系与定义条件。',
            'pages':[{'source_id':p['source_id'],
                'narration':'任务是要完成的工作，性能用来衡量工作的表现。经验是学习时获得的资料。性能随经验提高，是这个定义的核心条件。',
                'bullets':['性能随经验提高'],
                'visual':{'representation':'comparison','items':[
                    {'label':'任务','caption':'要完成的工作','sentence_id':1},
                    {'label':'性能','caption':'衡量完成的表现','sentence_id':1},
                    {'label':'经验','caption':'学习获得的资料','sentence_id':2}], 'links':[]},'asset_ids':[]}
                for p in data['points']]})


def test_complete_context_no_model_veto_all_points_and_cache(tmp_path):
    document,bundle=inputs();client=DraftClient()
    options=dict(template_id='editorial',use_illustrations=False,folder=tmp_path/'draft',progress=lambda *args:None)
    lesson=build_draft_lesson(client,bundle,document,'',VoiceMode.SYSTEM,**options)
    assert len(client.calls)==3 and len(lesson.segments)==5
    assert lesson.animation_report['content_review']=='not_run'
    assert not any('approved' in v for v in lesson.deck_plan['visual_story']['visuals'].values())
    assert lesson.deck_plan['front_matter'][0]=='cover' and lesson.deck_plan['back_matter']==['recap']
    validate_lesson(document,lesson)
    validate_choices([DeckPageChoice.model_validate(p) for p in lesson.deck_plan['pages']],lesson,
                     get_template('editorial'),range(1,6))
    again=build_draft_lesson(client,bundle,document,'',VoiceMode.SYSTEM,**options)
    assert len(client.calls)==3 and again.model_dump()==lesson.model_dump()


def test_invalid_assets_and_links_warn_omission_keeps_source(tmp_path):
    document,bundle=inputs(2)
    class Faulty(DraftClient):
        def generate(self,*args):
            draft=super().generate(*args)
            if not hasattr(draft,'pages'):return draft
            draft.pages=draft.pages[:1]
            draft.pages[0].asset_ids=['untrusted/path.svg']
            from zhijiang.slide_story import VisualLink
            draft.pages[0].visual.representation='process'
            draft.pages[0].visual.links=[VisualLink(source=1,target=4,label='参考',sentence_id=1)]
            return draft
    lesson=build_draft_lesson(Faulty(),bundle,document,'',VoiceMode.SYSTEM,template_id='paper',
        use_illustrations=False,folder=tmp_path/'draft',progress=lambda *args:None)
    assert len(lesson.segments)==2
    assert len(lesson.animation_report['warnings'])==3
    assert lesson.segments[1].evidence==bundle.points[1].evidence
    assert lesson.deck_plan['pages'][0]['asset_refs']==[]
    assert not lesson.deck_plan['visual_story']['visuals']['1']['visual']['links']


def test_shared_ppt_renders_unreviewed_draft_without_semantic_veto(tmp_path):
    document,bundle=inputs(2)
    lesson=build_draft_lesson(DraftClient(),bundle,document,'',VoiceMode.SYSTEM,template_id='editorial',
        use_illustrations=False,folder=tmp_path/'draft',progress=lambda *args:None)
    # Paraphrases need not be literal substrings. Structural rendering still runs.
    lesson.deck_plan['visual_story']['visuals']['1']['visual']['items'][0]['caption']='待完成的目标'
    # A producer's imperfect theme wording must not become another veto.
    lesson.deck_plan['visual_story']['navigation']['groups'][0]['label']='关于'
    from zhijiang.presentation import render_presentation
    from pptx import Presentation
    render_presentation(lesson,tmp_path/'lesson.pptx')
    deck=Presentation(tmp_path/'lesson.pptx')
    assert len(deck.slides)>=5
    assert any('经验是学习时获得' in s.notes_slide.notes_text_frame.text for s in deck.slides)
    assert list((tmp_path/'presentation-preview').glob('*.png'))


def test_default_pipeline_exports_real_ppt_and_mp4(tmp_path,sample_pdf):
    import wave
    from zhijiang.agents import AIAgents, DemoAgents
    from zhijiang.pdf import read_pdf
    from zhijiang.models import GenerationOptions, JobStatus, Mode
    from zhijiang.config import Settings
    from zhijiang.storage import JobStore
    from zhijiang.pipeline import JobProcessor
    document=read_pdf(sample_pdf,'sample.pdf')
    bundle=DemoAgents().extract_knowledge(document)
    class Model:
        model='mock';base_url='http://localhost';semantic_thinking=False
        def close(self):pass
        def generate(self,schema,instruction,material):
            data=json.loads(material)
            if schema.__name__=='DraftBookends':
                return schema.model_validate({'cut_after':[], 'theme_labels':['二分查找'],
                    'closing_ids':[data['closing_candidates'][0]['id']]})
            assert schema.__name__=='TeachingDraftBatch'
            return schema.model_validate({'title':'二分查找的条件','objective':'理解有序资料中逐步缩小查找范围的方法。',
                'pages':[{'source_id':p['source_id'],'narration':'理解二分查找之前，先注意资料需要有序。每次比较中间对象，然后缩小范围。这里展示的是教学关系图。',
                    'bullets':['有序资料逐步缩小范围'],
                    'visual':{'representation':'process','items':[
                        {'label':'有序资料','caption':'定位中间对象','sentence_id':1},
                        {'label':'查找范围','caption':'根据比较缩小范围','sentence_id':2}],
                        'links':[{'source':1,'target':2,'label':'比较','sentence_id':2,'kind':'directed'}]},'asset_ids':[]}
                    for p in data['points']]})
    class Agents(AIAgents):
        def extract_knowledge(self,*args,**kwargs):return bundle
    class Speech:
        def synthesize(self,text,output):
            with wave.open(str(output),'wb') as f:
                f.setnchannels(1);f.setsampwidth(2);f.setframerate(24000);f.writeframes(b'\0\0'*2400)
    from zhijiang.speech import SpeechError
    class OfflineVoice:
        def synthesize(self,text,output):raise SpeechError('服务未启动')
    store=JobStore(tmp_path/'data')
    job=store.create('sample.pdf',Mode.AI,VoiceMode.AI,True,False,sample_pdf,
        GenerationOptions(provider='ollama',base_url='http://localhost',model='mock',ppt_template='paper'))
    processor=JobProcessor(Settings(data_dir=store.data_dir),store,
        agent_factory=lambda _:Agents(Model()),speech_factory=lambda mode:Speech() if mode==VoiceMode.SYSTEM else OfflineVoice())
    processor.process(job['id'])
    state=store.get(job['id']);assert state['status']==JobStatus.COMPLETED,state['error']
    lesson=store.lesson(job['id'])
    assert lesson.voice_mode==VoiceMode.SYSTEM
    assert lesson.animation_report['warnings'][-1]['stage']=='speech'
    folder=store.jobs_dir/job['id']
    assert (folder/'lesson.pptx').stat().st_size>1000
    import imageio_ffmpeg
    frames=imageio_ffmpeg.read_frames(str(folder/'lesson.mp4'))
    info=next(frames);assert info['duration']>=8 and info['size']==(1280,720)
    assert next(frames);frames.close()


def test_bookends_choose_core_sentences_instead_of_final_metadata(tmp_path):
    from zhijiang.draft_bookends import enrich_draft_bookends
    document,bundle=inputs(4)
    lesson=build_draft_lesson(DraftClient(),bundle,document,'',VoiceMode.SYSTEM,
        template_id='paper',use_illustrations=False,folder=tmp_path/'initial',progress=lambda *args:None)
    lesson.segments[-1].title='相关期刊与组织'
    lesson.segments[-1].narration='教材列举了相关研究期刊与专业组织。它们可以作为寻找延伸资料的入口。'
    class BookendClient:
        model='test';base_url='http://localhost';semantic_thinking=False
        def generate(self,schema,instruction,material):
            data=json.loads(material)
            assert '不能把课程最后三段自动当总结' in instruction
            assert any(c['segment_id']==4 for c in data['closing_candidates'])
            chosen=next(c for c in data['closing_candidates']
                        if c['segment_id']==1 and '核心条件' in c['text'])
            return schema.model_validate({'cut_after':[3], 'theme_labels':['学习的条件','延伸资料'],
                'closing_ids':[chosen['id']]})
    enrich_draft_bookends(lesson,BookendClient(),folder=tmp_path/'bookends')
    story=lesson.deck_plan['visual_story']
    assert story['closing_points'][0]['text']=='性能随经验提高，是这个定义的核心条件'
    assert story['navigation']['takeaway_ids']==[1]
    assert [i for group in story['navigation']['groups'] for i in group['segment_ids']]==[1,2,3,4]
    assert story['content_review']=='not_run'


def test_bookend_service_failure_warns_without_rejecting_course(tmp_path):
    from zhijiang.agents import GenerationError
    from zhijiang.draft_bookends import enrich_draft_bookends
    document,bundle=inputs(2)
    lesson=build_draft_lesson(DraftClient(),bundle,document,'',VoiceMode.SYSTEM,
        template_id='paper',use_illustrations=False,folder=tmp_path/'initial',progress=lambda *args:None)
    original=json.dumps(lesson.deck_plan,sort_keys=True)
    class Offline:
        def generate(self,*args):raise GenerationError('模型服务断开')
    enrich_draft_bookends(lesson,Offline(),folder=tmp_path/'offline')
    assert json.dumps(lesson.deck_plan,sort_keys=True)==original
    assert lesson.animation_report['warnings'][-1]['stage']=='bookends'
