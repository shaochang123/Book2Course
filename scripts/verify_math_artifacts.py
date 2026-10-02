"""Verify a completed math job; decoding is not a human playback review.

Usage: python scripts/verify_math_artifacts.py --job-dir data/jobs/JOB_ID
Requires the optional math-animation dependencies.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

import av
from PIL import Image, ImageDraw
from pptx import Presentation

from zhijiang.agents import validate_evidence
from zhijiang.math_planning import verify_scene
from zhijiang.models import Lesson
from zhijiang.pdf import read_pdf


def media_check(path: Path) -> dict:
    with av.open(str(path)) as media:
        video = media.streams.video[0]
        audio = media.streams.audio[0]
        video_seconds = float(video.duration*video.time_base)
        audio_seconds = float(audio.duration*audio.time_base)
        details = {"video_codec": video.codec_context.name, "audio_codec": audio.codec_context.name,
                   "width": video.width, "height": video.height, "fps": float(video.average_rate),
                   "video_seconds": video_seconds, "audio_seconds": audio_seconds}
        assert (video.width, video.height) == (1280, 720), details
        assert float(video.average_rate) == 30, details
        assert abs(video_seconds-audio_seconds) < 0.12, details
        assert details["video_codec"] == "h264" and details["audio_codec"] == "aac", details
        frames = sum(1 for _ in media.decode(video))
        assert frames > 0
        details["decoded_video_frames"] = frames
    with av.open(str(path)) as media:
        samples, maximum = 0, 0.0
        for frame in media.decode(audio=0):
            values = frame.to_ndarray()
            maximum = max(maximum, float(abs(values).max()))
            samples += frame.samples
        assert samples > 0 and maximum > 0.001, "Missing or silent narration"
        details.update(decoded_audio_samples=samples, maximum_audio_amplitude=maximum)
    return details


def frame_at(path: Path, seconds: float) -> Image.Image:
    with av.open(str(path)) as media:
        stream = media.streams.video[0]
        media.seek(int(seconds/stream.time_base), stream=stream)
        for frame in media.decode(stream):
            if float(frame.pts*stream.time_base) >= seconds:
                return frame.to_image()
    raise ValueError(f"No frame at {seconds}")


def verify(folder: Path) -> dict:
    data = json.loads((folder/'math-scenes.json').read_text(encoding='utf-8'))
    lesson = Lesson.model_validate(data['lesson'])
    document = read_pdf((folder/'source.pdf').read_bytes(), 'source.pdf')
    assets = data['assets']
    assert len(assets) == len(lesson.segments) >= 3
    report = {"job_dir": str(folder), "source_sha256": hashlib.sha256((folder/'source.pdf').read_bytes()).hexdigest(),
              "scene_count": len(assets), "scenes": [], "human_review": {
                  "source_meaning": "not recorded by this script", "visual_layout": "not recorded by this script",
                  "full_playback_and_listening": "not performed by decoding", "powerpoint_slideshow": "not performed by OOXML checks"}}
    contact = Image.new('RGB',(960,270*len(assets)), '#081623')
    for index, (segment, asset) in enumerate(zip(lesson.segments, assets)):
        scene = segment.math_scene
        assert scene and scene.narration_binding == 'verified'
        algebra = verify_scene(scene)
        validate_evidence(document, scene.evidence)
        for citation in lesson.animation_report['sources'][scene.kind]:
            from zhijiang.models import Evidence
            validate_evidence(document, Evidence.model_validate(citation))
        path = folder/'math'/f'scene-{index+1:02d}'
        geometry = asset['rendered_geometry']
        assert [entry['action'] for entry in geometry] == [beat.action for beat in scene.beats]
        paths = [movement for entry in geometry for movement in entry.get('continuous_paths', [])]
        for movement in paths:
            assert movement['passed'] and len(movement['samples']) >= 3, movement
        if scene.kind == 'composition':
            assert len(paths) == 2 and all(item['operation'] == 'angle_interpolation' for item in paths)
        if scene.kind == 'projection':
            assert len(paths) == 1 and paths[0]['operation'] == 'normal_projection'
        for entry, timing in zip(geometry, asset['timing']):
            assert timing['start'] <= entry['rendered_seconds'] <= timing['start']+timing['duration']+1/30
            assert all(check['passed'] for check in entry['checks'])
        media = media_check(path/'clip.mp4')
        assert abs(media['video_seconds']-asset['seconds']) < 0.12
        # Review source, intermediate, and final states on a contact sheet.
        times = [0.6, asset['timing'][len(asset['timing'])//2]['start']+1,
                 max(0.6, asset['seconds']-1)]
        for column, seconds in enumerate(times):
            image = frame_at(path/'clip.mp4', seconds).resize((480,270))
            image.thumbnail((320,180))
            contact.paste(image,(320*column,270*index))
        ImageDraw.Draw(contact).text((12,270*index+190), f"{index+1}. {scene.kind}", fill='white')
        report['scenes'].append({"kind": scene.kind, "source_page": scene.evidence.page,
            "calculation_checks": algebra['checks'], "rendered_paths": paths, "media": media})
    report['full_video'] = media_check(folder/'lesson.mp4')
    assert abs(report['full_video']['video_seconds']-sum(asset['seconds'] for asset in assets)) < 0.3
    deck = Presentation(folder/'lesson.pptx')
    assert len(deck.slides) == 1+2*len(assets)
    videos = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.shape_type == 16:  # MSO_SHAPE_TYPE.MEDIA
                assert abs(shape.width/shape.height-16/9) < 0.001
    with ZipFile(folder/'lesson.pptx') as archive:
        videos = [name for name in archive.namelist() if name.startswith('ppt/media/') and name.endswith('.mp4')]
        svgs = [name for name in archive.namelist() if name.startswith('ppt/media/') and name.endswith('.svg')]
        hashes = {hashlib.sha256(archive.read(name)).hexdigest() for name in videos}
        assert len(videos) == len(svgs) == len(assets)
        for index in range(len(assets)):
            digest = hashlib.sha256((folder/'math'/f'scene-{index+1:02d}'/'clip.mp4').read_bytes()).hexdigest()
            assert digest in hashes, "PPT does not contain the same lesson clip"
        assert len([name for name in archive.namelist() if name.startswith('ppt/notesSlides/notesSlide') and name.endswith('.xml')]) == len(deck.slides)
    report['pptx'] = {"slides": len(deck.slides), "embedded_videos": len(videos), "embedded_svgs": len(svgs),
                      "aspect_ratio": "16:9", "same_clip_bytes": True, "speaker_notes": True}
    report['passed'] = True
    contact.save(folder/'validation-contact.png')
    (folder/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--job-dir',type=Path,required=True)
    args = parser.parse_args()
    result = verify(args.job_dir.resolve())
    print(json.dumps({key: result[key] for key in ('passed','scene_count','pptx')},ensure_ascii=False))
