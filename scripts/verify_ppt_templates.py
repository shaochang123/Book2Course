"""Compare built-in templates using the SAME existing lesson and media.

This checks layout planning and PPT packaging, not fresh textbook understanding
or mathematical correctness. Source lesson and figures are never authored here.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import imageio_ffmpeg
from pptx import Presentation
from zhijiang.agents import OllamaClient
from zhijiang.config import Settings
from zhijiang.deck_planning import plan_deck
from zhijiang.models import Lesson
from zhijiang.presentation import render_presentation


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--lesson', type=Path, required=True)
    parser.add_argument('--assets-dir', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--model', default='')
    parser.add_argument('--base-url', default='http://127.0.0.1:11434')
    parser.add_argument('--remote-consent', action='store_true')
    args = parser.parse_args()
    if args.model and Settings.model_is_remote('ollama', args.base_url, args.model) and not args.remote_consent:
        parser.error('远程模型会接收已有讲稿和来源引文；确认资料可发送后使用 --remote-consent。')
    source = args.lesson.read_bytes()
    lesson = Lesson.model_validate_json(source)
    args.output.mkdir(parents=True, exist_ok=True)
    client = OllamaClient(args.base_url, args.model, semantic_thinking=False,
                          prefer_json=True) if args.model else None
    results = []
    try:
        for template_id in ['paper', 'academic', 'editorial']:
            folder = args.output/template_id; folder.mkdir(exist_ok=True)
            for index, segment in enumerate(lesson.segments, 1):
                if segment.math_scene or segment.visual_scene:
                    kind = 'math' if segment.math_scene else 'visual'
                    origin = args.assets_dir/kind/f'scene-{index:02d}'
                    target = folder/kind/f'scene-{index:02d}'; target.mkdir(parents=True, exist_ok=True)
                    for name in ['summary.svg', 'summary.png', 'clip.mp4', 'poster.png']:
                        shutil.copyfile(origin/name, target/name)
                elif (args.assets_dir/f'audio-{index-1}.wav').is_file():
                    shutil.copyfile(args.assets_dir/f'audio-{index-1}.wav', folder/f'audio-{index-1}.wav')
            variant = lesson.model_copy(deep=True); variant.ppt_template = template_id
            variant.deck_plan = plan_deck(variant, template_id, client, output=folder/'presentation-plan.json')
            variant_json = variant.model_dump_json(indent=2)
            (folder/'lesson.json').write_text(variant_json, encoding='utf-8')
            path = folder/'lesson.pptx'; render_presentation(variant, path)
            with zipfile.ZipFile(path) as archive:
                assert archive.testzip() is None
                svgs = [name for name in archive.namelist() if name.endswith('.svg')]
                movies = [name for name in archive.namelist() if name.endswith('.mp4')]
                expected = sum(item['include_animation'] for item in variant.deck_plan['pages'])
                # ZIP may deduplicate byte-identical movies; count native media
                # shapes to verify requested per-segment animation pages.
                deck = Presentation(path)
                actual = sum('video' in shape._element.xml for slide in deck.slides for shape in slide.shapes)
                assert actual == expected and len(svgs) >= len(variant.segments)
                for i, name in enumerate(movies):
                    decoded = folder/f'check-movie-{i+1}.mp4'
                    decoded.write_bytes(archive.read(name))
                    process = subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), '-v', 'error',
                        '-i', str(decoded), '-f', 'null', '-'], capture_output=True, timeout=180)
                    if process.returncode:
                        raise RuntimeError('PPT 内嵌视频完整解码失败。')
                    decoded.unlink()
            result = {'template': template_id, 'planner': variant.deck_plan['planner'],
                      'source_lesson_sha256': hashlib.sha256(source).hexdigest(),
                      'pptx_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                      'slide_count': len(deck.slides), 'svg_count': len(svgs),
                      'video_count': len(movies), 'technical_checks_passed': True,
                      'fresh_textbook_generation': False, 'math_joint_acceptance': False,
                      'manual_visual_review': 'pending', 'powerpoint_playback': 'not_performed'}
            results.append(result)
            print(json.dumps(result, ensure_ascii=False), flush=True)
            (args.output/'results.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
    finally:
        if client:
            (args.output/'model-calls.json').write_text(json.dumps(client.call_metrics, ensure_ascii=False, indent=2), encoding='utf-8')
            if client.invalid_outputs:
                (args.output/'model-format-errors.json').write_text(json.dumps(client.invalid_outputs, ensure_ascii=False, indent=2), encoding='utf-8')
            client.close()


if __name__ == '__main__':
    main()
