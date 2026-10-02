"""Verify generic narrated scene graphs, preserving explicit review boundaries."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw
from pptx import Presentation

from zhijiang.models import Lesson
from zhijiang.agents import validate_evidence
from zhijiang.pdf import read_pdf
from zhijiang.visual_planning import verify_visual_scene, evaluate, CHECKERS, object_geometry, check_domain_state
try:
    from .verify_math_artifacts import media_check, frame_at, verify as verify_math
except ImportError:
    from verify_math_artifacts import media_check, frame_at, verify as verify_math


def check_moving_marker_pixels(scene, asset, clip):
    """Check real intermediate movie pixels, not only the renderer's data log.

    Unique-color dots are measurable; shared colors and other primitives remain
    covered by geometry logs and manual frames, without claiming pixel identity.
    """
    import numpy as np
    from zhijiang.visual_planning import states
    results=[]
    colors=[o.color.lower() for o in scene.objects]
    for i,(before,after,visible) in enumerate(states(scene)):
        if before==after: continue
        timing=asset['timing'][i];motion_frames=max(6,round(timing['duration']*30*.55))
        time=round(timing['start']*30)/30+motion_frames/60
        params={key:(before[key]+after[key])/2 for key in before}
        frame=np.asarray(frame_at(clip,time).convert('RGB')).astype(float)
        height,width=frame.shape[:2];scale=height/8
        for obj in scene.objects:
            if obj.kind!='dot' or obj.id not in visible or colors.count(obj.color.lower())!=1: continue
            old=object_geometry(obj,before)['points'][0];new=object_geometry(obj,after)['points'][0]
            if old==new: continue
            x,y=object_geometry(obj,params)['points'][0]
            px=width/2+scale*(-5.3+10.6*(x-scene.x_range[0])/(scene.x_range[1]-scene.x_range[0]))
            py=height/2-scale*(-1.8+4.3*(y-scene.y_range[0])/(scene.y_range[1]-scene.y_range[0]))
            color=np.array([int(obj.color[n:n+2],16) for n in (1,3,5)])
            mask=np.linalg.norm(frame-color,axis=2)<45
            # Exclude colored formula values and captions below the plot.
            mask[int(height/2+1.8*scale)+10:]=False
            ys,xs=np.where(mask); assert len(xs), f'No visible marker pixels: {obj.id}'
            distance=float(np.min(np.sqrt((xs-px)**2+(ys-py)**2)))
            assert distance<=12, f'Intermediate movie marker differs from geometry: {obj.id}, {distance:.2f}px'
            ox=width/2+scale*(-5.3+10.6*(old[0]-scene.x_range[0])/(scene.x_range[1]-scene.x_range[0]))
            oy=height/2-scale*(-1.8+4.3*(old[1]-scene.y_range[0])/(scene.y_range[1]-scene.y_range[0]))
            old_pixels=None
            if np.hypot(ox-px,oy-py)>35:
                old_pixels=int(np.sum((xs-ox)**2+(ys-oy)**2<12**2))
                assert old_pixels==0, f'Stale marker remains at its old position: {obj.id}, {old_pixels} pixels'
            results.append({'step':i+1,'object':obj.id,'time':time,'pixel_distance':distance,
                            'old_position_pixels':old_pixels,'passed':True})
    return results


def verify(folder: Path):
    if not (folder/'scene-data.json').is_file(): return verify_math(folder)
    data=json.loads((folder/'scene-data.json').read_text(encoding='utf-8'))
    lesson=Lesson.model_validate(data['lesson'])
    document=read_pdf((folder/'source.pdf').read_bytes(),'source.pdf')
    assets=data['assets']; assert len(assets)==len(lesson.segments)
    report={'passed':False,'scene_count':len(assets),'scenes':[],
        'human_review':{'source_meaning':'not performed by this script','layout':'not performed by decoding',
                        'listening':'not performed by decoding','powerpoint_slideshow':'not performed by OOXML inspection'}}
    contact=Image.new('RGB',(960,240*len(assets)),'#081623')
    hashes=set()
    for i,(segment,asset) in enumerate(zip(lesson.segments,assets)):
        scene=segment.visual_scene; assert scene
        validate_evidence(document,scene.evidence); computed=verify_visual_scene(scene)
        assert computed==asset['verification']
        path=folder/'visual'/f'scene-{i+1:02d}'
        assert len(asset['rendered_geometry'])==len(scene.beats)
        for j,rendered in enumerate(asset['rendered_geometry']):
            state=computed['states'][j]
            assert rendered['passed'] and all(c['passed'] and c['maximum_coordinate_error']<=1e-6 for c in rendered['checks'])
            assert rendered['geometry']=={o.id:object_geometry(o,state['parameters']) for o in scene.objects}
            samples=rendered['samples']; assert samples
            before=scene.parameters if j==0 else computed['states'][j-1]['parameters']
            if before!=state['parameters']:
                assert len(samples)>1 and samples[-1]['alpha']>0.999
            for sample in samples:
                check_domain_state(scene,sample['parameters'])
                for check in scene.checks:
                    assert CHECKERS[check.checker](evaluate(check.expression,sample['parameters']),check.expected,check.tolerance)
            assert abs(rendered['rendered_seconds']-(asset['timing'][j]['start']+asset['timing'][j]['duration']))<0.1
        media=media_check(path/'clip.mp4'); hashes.add(hashlib.sha256((path/'clip.mp4').read_bytes()).hexdigest())
        assert abs(media['video_seconds']-asset['seconds'])<0.12
        times=[.6,asset['timing'][len(scene.beats)//2]['start']+1,asset['seconds']-1]
        for col,t in enumerate(times): contact.paste(frame_at(path/'clip.mp4',t).resize((320,180)),(col*320,i*240))
        ImageDraw.Draw(contact).text((12,i*240+195),f'{i+1}. {scene.domain}',fill='white')
        pixels=check_moving_marker_pixels(scene,asset,path/'clip.mp4')
        report['scenes'].append({'domain':scene.domain,'source_page':scene.evidence.page,'checks':computed,'media':media,
                                 'intermediate_marker_pixels':pixels,
                                 'continuous_samples':sum(len(r['samples']) for r in asset['rendered_geometry'])})
    report['full_video']=media_check(folder/'lesson.mp4')
    deck=Presentation(folder/'lesson.pptx'); assert len(deck.slides)==1+2*len(assets)
    for index, asset in enumerate(assets):
        for slide_index in (1+2*index, 2+2*index):
            slide = deck.slides[slide_index]
            notes = slide.notes_slide.notes_text_frame.text
            for cue in asset['timing']:
                assert cue['text'] in notes, 'Speaker notes differ from synthesized step narration'
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.shape_type==16: assert abs(shape.width/shape.height-16/9)<0.001
    with ZipFile(folder/'lesson.pptx') as z:
        videos=[n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.mp4')]
        svgs=[n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.svg')]
        assert len(videos)==len(svgs)==len(assets)
        assert {hashlib.sha256(z.read(n)).hexdigest() for n in videos}==hashes
        assert len([n for n in z.namelist() if n.startswith('ppt/notesSlides/notesSlide') and n.endswith('.xml')])==len(deck.slides)
    report['pptx']={'slides':len(deck.slides),'embedded_videos':len(videos),'embedded_svgs':len(svgs),'same_clip_bytes':True,'aspect_ratio':'16:9','speaker_notes':True,'notes_match_spoken_steps':True}
    report['passed']=True; contact.save(folder/'validation-contact.png')
    (folder/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return report


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__); p.add_argument('--job-dir',type=Path,required=True)
    r=verify(p.parse_args().job_dir.resolve())
    print(json.dumps({k:r[k] for k in ['passed','scene_count','pptx']},ensure_ascii=False))
